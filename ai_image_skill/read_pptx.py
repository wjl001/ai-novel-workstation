# -*- coding: utf-8 -*-
from pptx import Presentation
import sys

prs = Presentation(r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\ai_image_skill\guide.pptx')
print(f'总幻灯片数: {len(prs.slides)}')
print()

for i, slide in enumerate(prs.slides):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t:
                    texts.append(t)
    if texts:
        print(f'--- 第{i+1}页 ---')
        for t in texts[:20]:
            print(f'  {t[:100]}')
        if len(texts) > 20:
            print(f'  ... (共{len(texts)}段)')
        print()
