# -*- coding: utf-8 -*-
"""验证Word文档：结构、字体、图片、表格、内容覆盖"""
from docx import Document
from docx.oxml.ns import qn
import os

DOC_PATH = r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\sensitive_word_v2\敏感词过滤引擎二期产品方案.docx'

doc = Document(DOC_PATH)

print('=' * 60)
print('文档验证报告')
print('=' * 60)

# 1. 基本统计
paras = doc.paragraphs
tables = doc.tables
print(f'\n【基本统计】')
print(f'  段落总数: {len(paras)}')
print(f'  表格总数: {len(tables)}')

# 统计图片
img_count = 0
for rel in doc.part.rels.values():
    if 'image' in rel.reltype:
        img_count += 1
print(f'  图片总数: {img_count}')

# 2. 字体检查
print(f'\n【字体检查】')
font_issues = []
checked = 0
for p in paras:
    for run in p.runs:
        if run.text.strip():
            checked += 1
            rPr = run._element.find(qn('w:rPr'))
            font_name = run.font.name
            east_asia = None
            if rPr is not None:
                rFonts = rPr.find(qn('w:rFonts'))
                if rFonts is not None:
                    east_asia = rFonts.get(qn('w:eastAsia'))
            if font_name != '微软雅黑' or east_asia != '微软雅黑':
                font_issues.append(f'  异常: text="{run.text[:20]}..." ascii={font_name} eastAsia={east_asia}')

print(f'  检查文本run数: {checked}')
print(f'  字体异常数: {len(font_issues)}')
for issue in font_issues[:5]:
    print(issue)
if len(font_issues) > 5:
    print(f'  ... 还有 {len(font_issues)-5} 条异常')

# 表格内字体检查
table_font_issues = 0
table_cell_count = 0
for table in tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    if run.text.strip():
                        table_cell_count += 1
                        if run.font.name != '微软雅黑':
                            table_font_issues += 1
print(f'  表格文本run数: {table_cell_count}')
print(f'  表格字体异常数: {table_font_issues}')

# 3. 内容覆盖检查
print(f'\n【内容覆盖检查】')
full_text = '\n'.join([p.text for p in paras])
for table in tables:
    for row in table.rows:
        for cell in row.cells:
            full_text += '\n' + cell.text

keywords = [
    '正向组合拦截', '反向语境豁免', '组合词匹配', '白名单',
    '产品架构', '业务流程', '泳道', '词库标签',
    '风险分级', '实施计划', '里程碑', '预期收益',
    '微软雅黑', '前缀+核心词', '误杀率', '漏检率',
    '领导汇报', '项目背景', '痛点分析',
]
for kw in keywords:
    found = kw in full_text
    status = '✓' if found else '✗ 缺失!'
    print(f'  {status} {kw}')

# 4. 章节结构检查
print(f'\n【章节结构检查】')
headings = []
for p in paras:
    text = p.text.strip()
    if text and (text.startswith('一、') or text.startswith('二、') or text.startswith('三、')
                  or text.startswith('四、') or text.startswith('五、') or text.startswith('六、')
                  or text.startswith('七、') or text.startswith('八、') or text.startswith('九、')):
        headings.append(text)
for h in headings:
    print(f'  {h}')

# 5. 表格内容检查
print(f'\n【表格检查】')
for i, table in enumerate(tables):
    rows = len(table.rows)
    cols = len(table.columns)
    first_cell = table.rows[0].cells[0].text[:20]
    print(f'  表{i+1}: {rows}行 x {cols}列, 首格="{first_cell}"')

print(f'\n【文件信息】')
print(f'  文件大小: {os.path.getsize(DOC_PATH) / 1024:.1f} KB')
print(f'  文件路径: {DOC_PATH}')

print('\n' + '=' * 60)
print('验证完成')
print('=' * 60)
