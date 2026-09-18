# -*- coding: utf-8 -*-
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
import os

fpath = r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\sensitive_word_v2\敏感词过滤引擎二期产品方案_V3.docx'
doc = Document(fpath)
paras = doc.paragraphs
tables = doc.tables

# 图片统计
img_count = 0
img_sizes = []
for rel in doc.part.rels.values():
    if 'image' in rel.reltype:
        img_count += 1

# 检查文档中图片的实际尺寸
for i, p in enumerate(paras):
    for run in p.runs:
        drawings = run._element.findall('.//' + qn('w:drawing'))
        for d in drawings:
            ext = d.find('.//' + qn('wp:extent'))
            if ext is not None:
                cx = int(ext.get('cx'))
                cy = int(ext.get('cy'))
                img_sizes.append((cx/914400, cy/914400))  # EMU to inches

print('=== 文档验证 V3 ===')
print(f'段落数: {len(paras)}')
print(f'表格数: {len(tables)}')
print(f'图片数(关系): {img_count}')
print(f'图片嵌入数(绘图): {len(img_sizes)}')
for i, (w, h) in enumerate(img_sizes):
    print(f'  图片{i+1}: {w:.2f} x {h:.2f} inches ({w*2.54:.1f} x {h*2.54:.1f} cm)')

# 字体检查
font_issues = 0
checked = 0
for p in paras:
    for r in p.runs:
        if r.text.strip():
            checked += 1
            if r.font.name != '微软雅黑':
                font_issues += 1
for t in tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    if r.text.strip():
                        checked += 1
                        if r.font.name != '微软雅黑':
                            font_issues += 1
print(f'字体检查: {checked}个文本run, 异常{font_issues}个')

# 内容关键词
full = '\n'.join(p.text for p in paras)
for t in tables:
    for row in t.rows:
        for cell in row.cells:
            full += '\n' + cell.text

keywords = ['三层过滤流水线','AC自动机','轻量语义','开源','大模型复核','阿里云','腾讯云','百度',
            '白名单','语境豁免','胸口射箭','狭小逼仄','干脆挂掉','前置审核','后置审核','Prompt',
            'Response','旁路','阻断','误杀率','漏判率','组合规则','归一化','形近字','谐音',
            '固定搭配','场景标签','词性分析','上下文窗口','熔断降级','text_scan','时序图','系统交互时序']
print('=== 内容覆盖 ===')
missing = []
for kw in keywords:
    ok = kw in full
    if not ok:
        missing.append(kw)
print(f'关键词: {len(keywords)}个, 覆盖{len(keywords)-len(missing)}个, 缺失{len(missing)}个')
if missing:
    print(f'缺失: {missing}')

# 图注检查
captions = [p.text for p in paras if p.text.startswith('图')]
print('=== 图注 ===')
for c in captions:
    print(f'  {c}')

# 一级章节
print('=== 一级章节 ===')
for p in paras:
    t = p.text.strip()
    if t and any(t.startswith(x) for x in ['一、','二、','三、','四、','五、','六、','七、','八、','九、']):
        print(f'  {t}')

print(f'文件大小: {os.path.getsize(fpath)/1024:.1f} KB')
print('=== 验证完成 ===')
