# -*- coding: utf-8 -*-
"""深度验证：图片嵌入、页面设置、标题层级"""
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Emu
import os

DOC_PATH = r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\sensitive_word_v2\敏感词过滤引擎二期产品方案.docx'
doc = Document(DOC_PATH)

print('=' * 60)
print('深度格式验证')
print('=' * 60)

# 1. 页面设置
print('\n【页面设置】')
section = doc.sections[0]
print(f'  页面宽度: {section.page_width / 360000:.1f} cm (标准A4=21.0)')
print(f'  页面高度: {section.page_height / 360000:.1f} cm (标准A4=29.7)')
print(f'  上边距: {section.top_margin / 360000:.1f} cm')
print(f'  下边距: {section.bottom_margin / 360000:.1f} cm')
print(f'  左边距: {section.left_margin / 360000:.1f} cm')
print(f'  右边距: {section.right_margin / 360000:.1f} cm')

# 2. 图片嵌入验证
print('\n【图片嵌入验证】')
img_paras = []
for i, p in enumerate(doc.paragraphs):
    for run in p.runs:
        if run._element.findall('.//' + qn('w:drawing')):
            drawings = run._element.findall('.//' + qn('w:drawing'))
            for d in drawings:
                ext = d.find('.//' + qn('wp:extent'))
                if ext is not None:
                    cx = int(ext.get('cx'))
                    cy = int(ext.get('cy'))
                    width_cm = cx / 360000
                    height_cm = cy / 360000
                    img_paras.append((i, width_cm, height_cm, p.alignment))

for idx, (para_idx, w, h, align) in enumerate(img_paras):
    align_str = '居中' if align == 1 else str(align)
    print(f'  图片{idx+1}: 段落#{para_idx}, 尺寸={w:.1f}x{h:.1f}cm, 对齐={align_str}')

# 3. 标题层级验证
print('\n【标题层级验证】')
h1_count = 0
h2_count = 0
h3_count = 0
for p in doc.paragraphs:
    text = p.text.strip()
    if not text:
        continue
    # 检查是否有左边框（一级标题特征）
    pPr = p._element.find(qn('w:pPr'))
    has_left_border = False
    if pPr is not None:
        pBdr = pPr.find(qn('w:pBdr'))
        if pBdr is not None:
            has_left_border = True
    # 检查字体大小
    for run in p.runs:
        if run.text.strip() and run.font.size:
            size = run.font.size.pt
            if size >= 15 and has_left_border:
                h1_count += 1
                print(f'  [H1] {text[:40]} (size={size}pt, 左边框=有)')
            elif size >= 12.5 and size < 15:
                h2_count += 1
            elif size >= 11 and size < 12.5 and run.font.bold:
                h3_count += 1
            break

print(f'\n  一级标题(H1): {h1_count}个')
print(f'  二级标题(H2): {h2_count}个')
print(f'  三级标题(H3): {h3_count}个')

# 4. 表格样式验证
print('\n【表格样式验证】')
for i, table in enumerate(doc.tables):
    # 检查表头背景色
    header_cell = table.rows[0].cells[0]
    tcPr = header_cell._tc.find(qn('w:tcPr'))
    shd = None
    if tcPr is not None:
        shd = tcPr.find(qn('w:shd'))
    fill = shd.get(qn('w:fill')) if shd is not None else '无'
    # 检查表头文字颜色
    header_run = None
    for p in header_cell.paragraphs:
        for r in p.runs:
            if r.text.strip():
                header_run = r
                break
        if header_run:
            break
    header_color = header_run.font.color.rgb if header_run and header_run.font.color and header_run.font.color.rgb else '默认'
    print(f'  表{i+1}: 表头背景={fill}, 表头字色={header_color}, 行数={len(table.rows)}')

# 5. 分页符检查
print('\n【分页符检查】')
page_breaks = 0
for p in doc.paragraphs:
    for run in p.runs:
        if run._element.findall(qn('w:br')):
            for br in run._element.findall(qn('w:br')):
                if br.get(qn('w:type')) == 'page':
                    page_breaks += 1
print(f'  分页符数量: {page_breaks}个 (预期9个，每章后分页)')

# 6. 字数统计
print('\n【字数统计】')
total_chars = 0
for p in doc.paragraphs:
    total_chars += len(p.text)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            total_chars += len(cell.text)
print(f'  全文字符数(含表格): {total_chars}')
print(f'  预估正文字数: 约{total_chars}字')

print('\n' + '=' * 60)
print('深度验证完成')
print('=' * 60)
