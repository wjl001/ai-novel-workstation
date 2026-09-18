# -*- coding: utf-8 -*-
from docx import Document
import os

doc = Document(r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\sensitive_word_v2\敏感词过滤引擎二期产品方案_V2.docx')
paras = doc.paragraphs
tables = doc.tables
img_count = sum(1 for rel in doc.part.rels.values() if 'image' in rel.reltype)

print('=== 文档验证 ===')
print(f'段落数: {len(paras)}')
print(f'表格数: {len(tables)}')
print(f'图片数: {img_count}')

font_issues = 0
checked = 0
for p in paras:
    for r in p.runs:
        if r.text.strip():
            checked += 1
            if r.font.name != '微软雅黑':
                font_issues += 1
for t in tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    if r.text.strip():
                        checked += 1
                        if r.font.name != '微软雅黑':
                            font_issues += 1
print(f'字体检查: {checked}个文本run, 异常{font_issues}个')

full = '\n'.join(p.text for p in paras)
for t in tables:
    for row in t.rows:
        for cell in row.cells:
            full += '\n' + cell.text

keywords = ['三层过滤流水线','AC自动机','轻量语义','开源','大模型复核','阿里云','腾讯云','百度',
            '白名单','语境豁免','胸口射箭','狭小逼仄','干脆挂掉','前置审核','后置审核','Prompt',
            'Response','旁路','阻断','误杀率','漏判率','组合规则','归一化','形近字','谐音',
            '固定搭配','场景标签','词性分析','上下文窗口','熔断降级','text_scan']
print('=== 内容覆盖 ===')
missing = []
for kw in keywords:
    ok = kw in full
    print(f'  {"OK" if ok else "MISS"} {kw}')
    if not ok:
        missing.append(kw)
print(f'缺失关键词: {missing if missing else "无"}')

print('=== 一级章节 ===')
for p in paras:
    t = p.text.strip()
    if t and any(t.startswith(x) for x in ['一、','二、','三、','四、','五、','六、','七、','八、','九、']):
        print(f'  {t}')

fpath = r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\sensitive_word_v2\敏感词过滤引擎二期产品方案_V2.docx'
print(f'文件大小: {os.path.getsize(fpath)/1024:.1f} KB')
