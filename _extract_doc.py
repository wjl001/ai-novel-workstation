import docx

doc = docx.Document(r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\视频生成Prompt智能增强系统_HermesAgent多轮对话版.docx')

with open(r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\_doc_full.txt', 'w', encoding='utf-8') as f:
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            f.write(f'[{i}] [{para.style.name}] {para.text}\n')
    for ti, table in enumerate(doc.tables):
        f.write(f'\n=== TABLE {ti} ===\n')
        for ri, row in enumerate(table.rows):
            cells = [cell.text.strip().replace(chr(10), ' ') for cell in row.cells]
            f.write(f'  ROW{ri}: ' + ' | '.join(cells) + '\n')

print('done, total paragraphs:', len(doc.paragraphs), 'tables:', len(doc.tables))
