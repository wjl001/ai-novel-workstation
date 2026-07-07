#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Markdown to DOCX Converter
"""

import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def set_run_font(run, font_name='宋体', font_size=10.5, bold=False):
    """设置字体"""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

def add_heading_zh(doc, text, level=1):
    """添加标题"""
    p = doc.add_paragraph()
    run = p.add_run(text)
    
    if level == 1:
        set_run_font(run, '黑体', 16, True)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 2:
        set_run_font(run, '黑体', 14, True)
    elif level == 3:
        set_run_font(run, '黑体', 12, True)
    else:
        set_run_font(run, '黑体', 11, True)
    
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)

def add_paragraph_zh(doc, text, bold=False, font_size=10.5, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    """添加段落"""
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, '宋体', font_size, bold)
    p.alignment = alignment
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    return p

def process_table(doc, lines, start_idx):
    """处理表格"""
    table_lines = []
    idx = start_idx
    
    while idx < len(lines) and lines[idx].strip().startswith('|'):
        table_lines.append(lines[idx])
        idx += 1
    
    if len(table_lines) < 2:
        return idx
    
    # 解析表头
    header_line = table_lines[0]
    headers = [cell.strip() for cell in header_line.split('|')[1:-1]]
    
    # 解析数据行
    data_rows = []
    for i in range(2, len(table_lines)):
        cells = [cell.strip() for cell in table_lines[i].split('|')[1:-1]]
        if cells:
            data_rows.append(cells)
    
    # 创建表格
    if headers:
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        
        # 填充表头
        hdr_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            hdr_cells[i].text = header
            for paragraph in hdr_cells[i].paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, '黑体', 10, True)
        
        # 填充数据行
        for row_data in data_rows:
            row_cells = table.add_row().cells
            for i, cell_text in enumerate(row_data):
                if i < len(row_cells):
                    row_cells[i].text = cell_text
                    for paragraph in row_cells[i].paragraphs:
                        for run in paragraph.runs:
                            set_run_font(run, '宋体', 10)
    
    doc.add_paragraph()
    return idx

def markdown_to_docx(input_file, output_file):
    """主转换函数"""
    
    # 读取文件
    with open(input_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # 创建文档
    doc = Document()
    
    # 设置默认样式
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style.font.size = Pt(10.5)
    
    # 添加封面
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run('AI短剧平台用户操作说明书')
    set_run_font(title_run, '黑体', 22, True)
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    info_para = doc.add_paragraph()
    info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_run = info_para.add_run('软件名称：AI短剧创作平台\n')
    set_run_font(info_run, '宋体', 12)
    info_run2 = info_para.add_run('开发单位：北京君禾世纪科技有限公司\n')
    set_run_font(info_run2, '宋体', 12)
    info_run3 = info_para.add_run('版本号：V1.0\n')
    set_run_font(info_run3, '宋体', 12)
    info_run4 = info_para.add_run('编写日期：2026年7月')
    set_run_font(info_run4, '宋体', 12)
    
    doc.add_page_break()
    
    # 处理正文
    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_content = []
    
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # 代码块处理
        if stripped.startswith('```'):
            if in_code_block:
                code_text = '\n'.join(code_content)
                add_paragraph_zh(doc, code_text, False, 9)
                code_content = []
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue
        
        if in_code_block:
            code_content.append(line)
            i += 1
            continue
        
        # 标题处理
        if stripped.startswith('# ') and not stripped.startswith('## '):
            add_heading_zh(doc, stripped[2:].strip(), 1)
            i += 1
            continue
        elif stripped.startswith('## '):
            add_heading_zh(doc, stripped[3:].strip(), 2)
            i += 1
            continue
        elif stripped.startswith('### '):
            add_heading_zh(doc, stripped[4:].strip(), 3)
            i += 1
            continue
        elif stripped.startswith('#### '):
            add_heading_zh(doc, stripped[5:].strip(), 4)
            i += 1
            continue
        
        # 表格处理
        if stripped.startswith('|') and i + 1 < len(lines) and lines[i + 1].strip().startswith('|'):
            i = process_table(doc, lines, i)
            continue
        
        # 列表处理
        if stripped.startswith('- ') or stripped.startswith('* '):
            item_text = stripped[2:].strip()
            item_text = re.sub(r'\*\*(.+?)\*\*', r'\1', item_text)
            p = add_paragraph_zh(doc, '• ' + item_text, False, 10.5)
            p.paragraph_format.left_indent = Inches(0.25)
            i += 1
            continue
        
        if re.match(r'^\d+\.\s', stripped):
            match = re.match(r'^(\d+)\.\s+(.+)$', stripped)
            if match:
                num = match.group(1)
                text = match.group(2)
                text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
                p = add_paragraph_zh(doc, num + '. ' + text, False, 10.5)
                p.paragraph_format.left_indent = Inches(0.25)
            i += 1
            continue
        
        # 普通段落
        if stripped:
            bold_pattern = re.compile(r'\*\*(.+?)\*\*')
            parts = bold_pattern.split(stripped)
            is_bold = False
            
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)
            
            for part in parts:
                if part:
                    run = p.add_run(part)
                    set_run_font(run, '宋体', 10.5, is_bold)
                is_bold = not is_bold
        else:
            doc.add_paragraph()
        
        i += 1
    
    # 保存文档
    doc.save(output_file)
    print(f"转换完成！已保存至: {output_file}")

if __name__ == '__main__':
    input_file = r'd:\phpstudy_pro\WWW\ai-novel-workstation2.7\docs\软著申请\AI短剧平台用户操作说明书.md'
    output_file = r'd:\phpstudy_pro\WWW\ai-novel-workstation2.7\docs\软著申请\AI短剧平台用户操作说明书.docx'
    
    markdown_to_docx(input_file, output_file)
