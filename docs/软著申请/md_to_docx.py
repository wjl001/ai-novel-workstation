#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Markdown to DOCX Converter for AI Short Drama Platform User Manual
"""

import sys
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

def set_chinese_font(run, font_name='宋体', font_size=10.5, bold=False):
    """设置中文字体"""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

def add_heading_zh(doc, text, level=1):
    """添加中文标题"""
    # 定义标题样式
    style_name = f'Heading {level}'
    
    # 添加段落
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles[style_name]
    
    # 添加文本
    run = paragraph.add_run(text)
    
    # 设置中文字体
    if level == 1:
        set_chinese_font(run, '黑体', 16, True)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 2:
        set_chinese_font(run, '黑体', 14, True)
    elif level == 3:
        set_chinese_font(run, '黑体', 12, True)
    else:
        set_chinese_font(run, '黑体', 11, True)
    
    return paragraph

def add_paragraph_zh(doc, text, bold=False, italic=False, font_size=10.5, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    """添加中文段落"""
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_chinese_font(run, '宋体', font_size, bold)
    run.font.italic = italic
    paragraph.alignment = alignment
    
    # 设置行距
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(6)
    
    return paragraph

def process_table(doc, lines, start_idx):
    """处理Markdown表格"""
    table_lines = []
    idx = start_idx
    
    # 收集表格行
    while idx < len(lines) and lines[idx].strip().startswith('|'):
        table_lines.append(lines[idx])
        idx += 1
    
    if len(table_lines) < 2:
        return idx, None
    
    # 解析表头
    header_line = table_lines[0]
    headers = [cell.strip() for cell in header_line.split('|')[1:-1]]
    
    # 解析数据行（跳过分隔行）
    data_rows = []
    for i in range(2, len(table_lines)):
        cells = [cell.strip() for cell in table_lines[i].split('|')[1:-1]]
        if cells:
            data_rows.append(cells)
    
    # 创建Word表格
    if headers:
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        
        # 填充表头
        hdr_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            hdr_cells[i].text = header
            # 设置表头样式
            for paragraph in hdr_cells[i].paragraphs:
                for run in paragraph.runs:
                    set_chinese_font(run, '黑体', 10, True)
        
        # 填充数据行
        for row_data in data_rows:
            row_cells = table.add_row().cells
            for i, cell_text in enumerate(row_data):
                if i < len(row_cells):
                    row_cells[i].text = cell_text
                    # 设置单元格样式
                    for paragraph in row_cells[i].paragraphs:
                        for run in paragraph.runs:
                            set_chinese_font(run, '宋体', 10)
    
    doc.add_paragraph()  # 表格后添加空行
    return idx, table

def markdown_to_docx(input_file, output_file):
    """将Markdown文件转换为DOCX文件"""
    
    # 读取Markdown文件
    with open(input_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # 创建Word文档
    doc = Document()
    
    # 设置文档默认字体
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style.font.size = Pt(10.5)
    
    # 添加封面
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run('AI短剧平台用户操作说明书')
    set_chinese_font(title_run, '黑体', 22, True)
    
    doc.add_paragraph()
    
    info_para = doc.add_paragraph()
    info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_run = info_para.add_run('软件名称：AI短剧创作平台\n')
    set_chinese_font(info_run, '宋体', 12)
    info_run2 = info_para.add_run('开发单位：北京君禾世纪科技有限公司\n')
    set_chinese_font(info_run2, '宋体', 12)
    info_run3 = info_para.add_run('版本号：V1.0\n')
    set_chinese_font(info_run3, '宋体', 12)
    info_run4 = info_para.add_run('编写日期：2026年7月')
    set_chinese_font(info_run4, '宋体', 12)
    
    # 添加分页符
    doc.add_page_break()
    
    # 处理内容
    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_content = []
    
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # 处理代码块
        if stripped.startswith('```'):
            if in_code_block:
                # 结束代码块
                code_text = '\n'.join(code_content)
                add_paragraph_zh(doc, code_text, False, False, 9, WD_ALIGN_PARAGRAPH.LEFT)
                code_content = []
                in_code_block = False
            else:
                # 开始代码块
                in_code_block = True
            i += 1
            continue
        
        if in_code_block:
            code_content.append(line)
            i += 1
            continue
        
        # 处理标题
        if stripped.startswith('# ') and not stripped.startswith('## '):
            # 一级标题
            title_text = stripped[2:].strip()
            add_heading_zh(doc, title_text, 1)
            i += 1
            continue
        elif stripped.startswith('## '):
            # 二级标题
            title_text = stripped[3:].strip()
            add_heading_zh(doc, title_text, 2)
            i += 1
            continue
        elif stripped.startswith('### '):
            # 三级标题
            title_text = stripped[4:].strip()
            add_heading_zh(doc, title_text, 3)
            i += 1
            continue
        elif stripped.startswith('#### '):
            # 四级标题
            title_text = stripped[5:].strip()
            add_heading_zh(doc, title_text, 4)
            i += 1
            continue
        
        # 处理表格
        if stripped.startswith('|') and i + 1 < len(lines) and lines[i + 1].strip().startswith('|'):
            i, table = process_table(doc, lines, i)
            continue
        
        # 处理列表
        if stripped.startswith('- ') or stripped.startswith('* '):
            item_text = stripped[2:].strip()
            # 处理加粗文本
            item_text = re.sub(r'\*\*(.+?)\*\*', r'\1', item_text)
            para = add_paragraph_zh(doc, '• ' + item_text, False, False, 10.5, WD_ALIGN_PARAGRAPH.LEFT)
            para.paragraph_format.left_indent = Inches(0.25)
            i += 1
            continue
        
        if re.match(r'^\d+\.\s', stripped):
            # 数字列表
            match = re.match(r'^(\d+)\.\s+(.+)$', stripped)
            if match:
                num = match.group(1)
                text = match.group(2)
                text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
                para = add_paragraph_zh(doc, num + '. ' + text, False, False, 10.5, WD_ALIGN_PARAGRAPH.LEFT)
                para.paragraph_format.left_indent = Inches(0.25)
            i += 1
            continue
        
        # 处理普通段落
        if stripped:
            # 处理加粗文本 **text**
            bold_pattern = re.compile(r'\*\*(.+?)\*\*')
            
            # 分段处理加粗文本
            parts = bold_pattern.split(stripped)
            is_bold = False
            
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_after = Pt(6)
            
            for part in parts:
                if part:
                    run = paragraph.add_run(part)
                    set_chinese_font(run, '宋体', 10.5, is_bold)
                is_bold = not is_bold
            
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            # 空行
            doc.add_paragraph()
        
        i += 1
    
    # 保存文档
    doc.save(output_file)
    print(f"转换完成！已保存至: {output_file}")

if __name__ == '__main__':
    input_file = r'd:\phpstudy_pro\WWW\ai-novel-workstation2.7\docs\软著申请\AI短剧平台用户操作说明书.md'
    output_file = r'd:\phpstudy_pro\WWW\ai-novel-workstation2.7\docs\软著申请\AI短剧平台用户操作说明书.docx'
    
    markdown_to_docx(input_file, output_file)
