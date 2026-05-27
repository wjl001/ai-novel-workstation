from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_ppt_master():
    prs = Presentation()
    # 采用宽屏比例 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 颜色配置
    color_primary = RGBColor(0, 82, 204)      # 主题蓝
    color_secondary = RGBColor(235, 87, 87)   # 强调红
    color_text_dark = RGBColor(51, 51, 51)    # 深灰文本
    color_text_light = RGBColor(102, 102, 102)# 浅灰文本
    color_bg = RGBColor(248, 249, 250)        # 极简灰白背景

    def apply_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color_bg

    # 1. 封面页
    slide_layout_title = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout_title)
    apply_background(slide)
    
    title = slide.shapes.title
    title.text = "我们的突围战：做AI短剧界的“超级工厂”"
    title.text_frame.paragraphs[0].font.name = "Microsoft YaHei"
    title.text_frame.paragraphs[0].font.size = Pt(54)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = color_primary

    subtitle = slide.placeholders[1]
    subtitle.text = "核心观点：不跟大厂抢流量，做包工头最爱的生产力工具。\n\n汇报主题：PC端AI短剧生产平台竞品调研与立项规划\n汇报人：[您的名字/部门]\n汇报日期：2026年5月"
    for p in subtitle.text_frame.paragraphs:
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(24)
        p.font.color.rgb = color_text_dark

    # 通用内容页函数
    def add_content_slide(title_text, core_point, bullet_points):
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        apply_background(slide)
        
        # 标题
        title = slide.shapes.title
        title.text = title_text
        title.text_frame.paragraphs[0].font.name = "Microsoft YaHei"
        title.text_frame.paragraphs[0].font.size = Pt(40)
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].font.color.rgb = color_primary
        title.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        
        # 内容区
        content = slide.placeholders[1]
        tf = content.text_frame
        
        # 核心观点
        p_core = tf.paragraphs[0]
        p_core.text = f"【核心观点】 {core_point}"
        p_core.font.name = "Microsoft YaHei"
        p_core.font.size = Pt(28)
        p_core.font.bold = True
        p_core.font.color.rgb = color_secondary
        p_core.space_after = Pt(20)
        
        # 详细要点
        for bp in bullet_points:
            p = tf.add_paragraph()
            p.text = bp
            p.level = 1
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(24)
            p.font.color.rgb = color_text_dark
            p.space_after = Pt(10)
            
    # 2. 市场现状
    add_content_slide(
        "[市场现状] 市场变天：横店没戏拍了，AI掀了桌子",
        "真人短剧正被淘汰，AI短剧已经占据绝对统治地位。",
        [
            "大盘数据：2026年Q1微短剧上线12.8万部，AI制作高达12.2万部。",
            "市场占比：AI短剧渗透率已超 95%！",
            "行业现状：传统剧组无戏可拍，算力（显卡）正在取代场地和群演。"
        ]
    )

    # 3. 核心驱动力
    add_content_slide(
        "[核心驱动力] 为什么大家都疯了？账算得太明白了",
        "成本砍掉90%，资本和制作团队必然全面倒向AI。",
        [
            "传统模式：场地+群演+灯光 = 50-100万成本，周期半个月。",
            "AI模式：3-5人+电脑+键盘 = 3-5万成本，周期1个星期。",
            "结论：资本是逐利的，十倍的利润差是市场剧变的唯一动力。"
        ]
    )

    # 4. 用户痛点
    add_content_slide(
        "[用户痛点] 工作室每天都在骂什么？",
        "虽然用了AI，但现在的工具根本没法干“重体力活”。",
        [
            "痛点一：“换头怪”灾难。角色脸部无法固定，第一集彭于晏，第十集吴彦祖，观众极易出戏。",
            "痛点二：“缝衣式”工作流。写剧本(GPT) -> 跑图(MJ) -> 生成视频(可灵) -> 剪辑(剪映)。",
            "现状总结：几千个镜头，切几十个文件夹，大半夜能把人逼疯。"
        ]
    )

    # 5. 竞品分析
    add_content_slide(
        "[竞品分析] 扒一扒对手：大厂们在干嘛？",
        "大厂都在拼底层大模型（造发动机），没空做体验（造车壳）。",
        [
            "字节（小云雀/即梦）：Seedance 2.0 模型，懂导演思维，镜头连贯。",
            "快手（可灵）：Kling 3.0，物理引擎牛，质感像电影大片。",
            "阿里（快乐小马）：便宜又快，自带翻译，主打出海管饱。"
        ]
    )

    # 6. 战略洞察
    add_content_slide(
        "[战略洞察] 为什么大厂不做“手机端短剧APP”？",
        "大厂绝不内部打架，工具端不能和流量端（抖音/快手）抢饭碗。",
        [
            "定位不同：AI工具是“炒菜的锅”，抖音快手是“大饭店”。",
            "流量闭环：大厂巴不得大家用他们的工具生成视频，再发回他们主站变现。",
            "物理限制：手机屏幕太小，搓几十集、几百个转场的长剧，完全不现实。"
        ]
    )

    # 7. 我们的定位
    add_content_slide(
        "[我们的定位] 避开大厂锋芒，做“包工头”最爱的软件",
        "我们不造发动机（大模型），我们造最好的全自动流水线机床！",
        [
            "绝对不碰：烧钱的底层大模型、C端短剧流量分发。",
            "死死咬住：B端工作室的“生产效率”和“工程化管理”。",
            "目标用户：急着赚钱、手里攥着剧本、被零碎软件折磨的短剧包工头。"
        ]
    )

    # 8. 杀手锏 1
    add_content_slide(
        "[杀手锏 1] 剧本一键拆解，把“作坊”变成“正规军”",
        "消灭乱七八糟的文件夹，建立正规的工程化目录树。",
        [
            "一键导入：十几万字剧本拖入，后台大模型瞬间解析。",
            "结构化拆解：自动生成 集数 -> 场次 -> 镜头分镜 -> 角色台词。",
            "左树右图：左侧点选镜头，右侧直接出画面生成面板。"
        ]
    )

    # 9. 杀手锏 2
    add_content_slide(
        "[杀手锏 2] “锁死”那张脸，打造短剧界“横店资产库”",
        "彻底解决“换头怪”痛点，让数字演员随叫随到。",
        [
            "开拍前捏脸：建立专属“数字演员休息室”。",
            "资产固化：死死锁住人脸特征、服装风格、AI配音音色。",
            "跨集调用：第一集到第三十集，一键“调用男主A”，角度随便换，脸绝对不崩。"
        ]
    )

    # 10. 杀手锏 3
    add_content_slide(
        "[杀手锏 3] “海纳百川”的智能调度器",
        "聚合各家API，用户只需提需求，我们负责跑腿。",
        [
            "痛点：买三家会员，开三个网页。",
            "解法：全面接入大厂开放API。要网感调字节，要特效调可灵，要便宜调阿里。",
            "本质：做一个“AI接口二道贩子”，用极致体验赚工作流的钱。"
        ]
    )

    # 11. 商业模式
    add_content_slide(
        "[商业模式] 怎么赚钱？（搞钱才是硬道理）",
        "细水长流收月租，赚差价收云租金，出海风口卖增值。",
        [
            "SaaS高级会员：卖剧本拆解、资产锁、无限轨剪辑等核心功能（省剪辑师人工费）。",
            "算力差价与云存储：拿大厂批发价赚差价；收几十G工程文件的云端保存费。",
            "“一键出海”增值包：按次收费，提供“一键翻译+口型同步”。"
        ]
    )

    # 12. 下一步行动
    add_content_slide(
        "[下一步行动] 我们的 Action Plan",
        "小步快跑，快速验证，死磕前端体验。",
        [
            "启动MVP开发：先做“剧本拆解成目录树” + “简单接入一个视频接口”。",
            "引入种子用户：找 2-3 家痛苦的短剧小团队入伙免费测，照着痛点改。",
            "前端资源倾斜：HomeView.vue 等前端交互必须像专业剪辑软件一样丝滑，拒绝网页卡顿感。"
        ]
    )

    # 13. 总结愿景
    add_content_slide(
        "[总结愿景] 卖铲子的人，永远最赚钱",
        "在AI短剧的淘金热中，我们就是那把最好用的冲锋枪。",
        [
            "旧规矩全被打破，影视工业面临重构。",
            "降维打击市面上的“网页玩具”。",
            "只要咱们这把“冲锋枪”造得好，必定能在这个风口赚得盆满钵满！"
        ]
    )

    # 14. 结尾页
    layout_end = prs.slide_layouts[0]
    slide_end = prs.slides.add_slide(layout_end)
    apply_background(slide_end)
    title_end = slide_end.shapes.title
    title_end.text = "Q & A"
    title_end.text_frame.paragraphs[0].font.name = "Microsoft YaHei"
    title_end.text_frame.paragraphs[0].font.size = Pt(72)
    title_end.text_frame.paragraphs[0].font.bold = True
    title_end.text_frame.paragraphs[0].font.color.rgb = color_primary

    subtitle_end = slide_end.placeholders[1]
    subtitle_end.text = "感谢聆听，请领导批评指正\n\n感谢各位领导的时间。\n欢迎提问与探讨。"
    for p in subtitle_end.text_frame.paragraphs:
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(28)
        p.font.color.rgb = color_text_dark

    prs.save('AI短剧平台立项规划-专业版.pptx')
    print("Master PPT generated successfully!")

if __name__ == '__main__':
    create_ppt_master()