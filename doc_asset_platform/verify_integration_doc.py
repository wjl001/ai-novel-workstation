# -*- coding: utf-8 -*-
from docx import Document
from docx.oxml.ns import qn
import os

DOC = r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\doc_asset_platform\新增厂商与新增模型接入流程及工作量评估.docx'
doc = Document(DOC)

print("=" * 55)
print("文档验证报告")
print("=" * 55)

print(f"\n【基本统计】")
print(f"  段落: {len(doc.paragraphs)}")
print(f"  表格: {len(doc.tables)}")
imgs = sum(1 for r in doc.part.rels.values() if 'image' in r.reltype)
print(f"  图片: {imgs}")

print(f"\n【字体检查】")
issues = []
yahei = 0
total = 0
for p in doc.paragraphs:
    for r in p.runs:
        total += 1
        rPr = r._element.find(qn('w:rPr'))
        if rPr is not None:
            rf = rPr.find(qn('w:rFonts'))
            if rf is not None:
                ea = rf.get(qn('w:eastAsia'))
                if ea and '雅黑' in ea:
                    yahei += 1
                elif ea:
                    issues.append(f"非雅黑: {ea}")
print(f"  正文run: {total}, 雅黑: {yahei}, 问题: {len(issues)}")

table_yahei = 0
table_total = 0
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    table_total += 1
                    rPr = r._element.find(qn('w:rPr'))
                    if rPr is not None:
                        rf = rPr.find(qn('w:rFonts'))
                        if rf is not None:
                            ea = rf.get(qn('w:eastAsia'))
                            if ea and '雅黑' in ea:
                                table_yahei += 1
print(f"  表格run: {table_total}, 雅黑: {table_yahei}")

print(f"\n【标题结构】")
for p in doc.paragraphs:
    if p.style.name.startswith('Heading'):
        print(f"  {p.style.name}: {p.text}")

print(f"\n【表格概览】")
for i, t in enumerate(doc.tables):
    print(f"  表{i+1}: {len(t.rows)}行x{len(t.columns)}列, 首格: {t.rows[0].cells[0].text[:15]}")

print(f"\n【文件大小】: {os.path.getsize(DOC)/1024:.1f} KB")
print("=" * 55)
