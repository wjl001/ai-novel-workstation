# -*- coding: utf-8 -*-
from docx import Document
from docx.oxml.ns import qn
import os

DOC = r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\AI润色优化_固定模板对比测试方案_字数差异化版_含固定模板.docx'
doc = Document(DOC)

print('=' * 60)
print('文档验证报告')
print('=' * 60)

print(f'\n【基本统计】')
print(f'  段落数: {len(doc.paragraphs)}')
print(f'  表格数: {len(doc.tables)}')
print(f'  文件大小: {os.path.getsize(DOC)/1024:.1f} KB')

# 检查新增内容
print(f'\n【新增内容检查】')
keywords = [
    'AI润色优化固定提示词模板',
    '点击按钮后核心处理逻辑',
    'AI润色按钮点击后处理逻辑',
    '角色四视图优化提示词模板',
    '场景优化提示词模板',
    '道具优化提示词模板',
    '角色四视图固定优化模板',
    '场景固定优化模板',
    '道具固定优化模板',
    '占位符说明',
    '三大模板对比与适用场景',
]
full_text = '\n'.join([p.text for p in doc.paragraphs])
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            full_text += '\n' + cell.text

for kw in keywords:
    found = kw in full_text
    print(f'  {"✓" if found else "✗"} {kw}')

# 字体检查
print(f'\n【字体检查】')
yahei_count = 0
total_runs = 0
font_issues = []
for p in doc.paragraphs:
    for r in p.runs:
        total_runs += 1
        rPr = r._element.find(qn('w:rPr'))
        if rPr is not None:
            rf = rPr.find(qn('w:rFonts'))
            if rf is not None:
                ea = rf.get(qn('w:eastAsia'))
                if ea and '雅黑' in ea:
                    yahei_count += 1
                elif ea:
                    font_issues.append(f'非雅黑: {ea}')

table_yahei = 0
table_total = 0
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    table_total += 1
                    rPr = r._element.find(qn('w:rPr'))
                    if rPr is not None:
                        rf = rPr.find(qn('w:rFonts'))
                        if rf is not None:
                            ea = rf.get(qn('w:eastAsia'))
                            if ea and '雅黑' in ea:
                                table_yahei += 1

print(f'  正文run: {total_runs}, 雅黑: {yahei_count}, 问题: {len(font_issues)}')
print(f'  表格run: {table_total}, 雅黑: {table_yahei}')

# 检查插入位置
print(f'\n【插入位置验证】')
for i, p in enumerate(doc.paragraphs):
    if '固定提示词模板' in p.text or '二、测试方法' in p.text:
        print(f'  [{i}] {p.text[:60]}')

# 检查模板表格内容
print(f'\n【模板表格内容抽样】')
for i, t in enumerate(doc.tables):
    first_text = t.rows[0].cells[0].text[:50]
    if '固定优化模板' in first_text or '占位符' in first_text or '处理逻辑' in first_text or '模板对比' in first_text:
        print(f'  表{i}: {first_text}')
        if t.rows[0].cells[0].text.strip():
            content = t.rows[1].cells[0].text[:120] if len(t.rows) > 1 else ''
            print(f'       内容预览: {content}')

print('\n' + '=' * 60)
print('验证完成')
print('=' * 60)
