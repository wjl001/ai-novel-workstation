# -*- coding: utf-8 -*-
"""验证Word文档"""
from docx import Document
import os

DOC_PATH = r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\doc_hermes\视频生成Prompt智能增强系统_汇报版.docx'

doc = Document(DOC_PATH)

print('=== 文档验证报告 ===')
print(f'文件大小: {os.path.getsize(DOC_PATH)} bytes')

# 统计段落
print(f'\n段落总数: {len(doc.paragraphs)}')
print(f'表格总数: {len(doc.tables)}')

# 统计图片
image_count = 0
for rel in doc.part.rels.values():
    if "image" in rel.reltype:
        image_count += 1
print(f'图片总数: {image_count}')

# 检查字体
print('\n=== 字体检查 ===')
fonts = set()
for para in doc.paragraphs:
    for run in para.runs:
        if run.font.name:
            fonts.add(run.font.name)
        if run.element.rPr is not None:
            rFonts = run.element.rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
            if rFonts is not None:
                ea = rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia')
                if ea:
                    fonts.add(f'eastAsia:{ea}')
print(f'使用的字体: {fonts}')

# 检查标题
print('\n=== 章节标题 ===')
for para in doc.paragraphs:
    text = para.text.strip()
    if text and (text.startswith('一、') or text.startswith('二、') or text.startswith('三、') or 
                 text.startswith('四、') or text.startswith('五、') or text.startswith('六、') or
                 text.startswith('七、') or text.startswith('八、') or text.startswith('九、') or
                 text.startswith('十、') or text.startswith('视频生成')):
        print(f'  - {text[:50]}')

# 检查表格
print('\n=== 表格内容检查 ===')
for i, table in enumerate(doc.tables):
    print(f'表格{i+1}: {len(table.rows)}行 x {len(table.columns)}列')
    print(f'  表头: {[cell.text[:15] for cell in table.rows[0].cells]}')

# 正文字数统计（不含表格）
body_chars = sum(len(p.text) for p in doc.paragraphs)
table_chars = 0
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            table_chars += len(cell.text)
print(f'\n=== 字数统计 ===')
print(f'正文字符数: {body_chars}')
print(f'表格字符数: {table_chars}')
print(f'总字符数: {body_chars + table_chars}')

# 检查图片引用
print('\n=== 图片引用检查 ===')
img_count_in_doc = 0
for para in doc.paragraphs:
    for run in para.runs:
        if run._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
            img_count_in_doc += 1
print(f'文档中引用的图片数: {img_count_in_doc}')

print('\n=== 验证完成 ===')
