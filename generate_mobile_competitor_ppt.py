import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def set_font(font, name='微软雅黑', size=None, bold=False, color=None):
    font.name = name
    if size:
        font.size = size
    font.bold = bold
    if color:
        font.color.rgb = color

def add_title_slide(prs, title_text, subtitle_text):
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = title_text
    subtitle.text = subtitle_text
    
    set_font(title.text_frame.paragraphs[0].font, size=Pt(44), bold=True)
    set_font(subtitle.text_frame.paragraphs[0].font, size=Pt(24))

def add_content_slide(prs, title_text, content_lines, conclusion_text=None):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = title_text
    set_font(title.text_frame.paragraphs[0].font, size=Pt(32), bold=True)
    
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    tf.clear()
    
    for line in content_lines:
        p = tf.add_paragraph()
        p.text = line
        set_font(p.font, size=Pt(18))
        p.level = 0
        p.space_after = Pt(10)
        
    if conclusion_text:
        p = tf.add_paragraph()
        p.text = "【总结结论】" + conclusion_text
        set_font(p.font, size=Pt(18), bold=True, color=RGBColor(255, 0, 0))
        p.level = 0
        p.space_before = Pt(14)
        
    # 添加占位截图框
    left = Inches(7.0)
    top = Inches(2.0)
    width = Inches(2.5)
    height = Inches(4.5)
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(230, 230, 230)
    text_frame = rect.text_frame
    text_frame.text = "移动端/产品截图占位"
    text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    set_font(text_frame.paragraphs[0].font, size=Pt(14), color=RGBColor(100, 100, 100))

def main():
    prs = Presentation()
    
    # 1. 封面
    add_title_slide(prs, "国内AI短剧平台移动端竞品调研", "小云雀、即梦、可灵、纳米、天工、万兴 竞品分析与差异化策略\n\n汇报人：产品团队\n日期：2026年5月")
    
    # 2. 市场综述
    add_content_slide(prs, "市场综述：AIGC短剧的移动端现状", [
        "核心发现：目前主流平台本质是AIGC生产端工具（锄头），而非短剧消费平台（粮仓）。",
        "PC端定位：复杂的长文本解析、多角色一致性、分镜管理和长线时间轴等“重度工作流”均依赖PC大屏。",
        "移动端定位：工具App的复杂操作会与移动端天然的“短平快”创作流冲突。",
        "生态流向：生成的内容天然流向抖音、快手等自家短视频平台变现，而非在生成工具内消费。"
    ], "多数平台未在移动端布局深度短剧创作，这为我们留下了“移动端轻量级创作+消费闭环”的差异化突围空间。")
    
    # 3. 竞品一：小云雀
    add_content_slide(prs, "竞品一：小云雀", [
        "PC端功能：公司已上线的短剧PC端平台，主打AIGC短剧完整生产流程。",
        "移动端有无：目前无专门针对短剧的移动端App。",
        "移动端功能：缺乏对应功能。",
        "PC与移动端关系：重度依赖PC，无移动端生态引流。"
    ], "小云雀目前纯聚焦PC端生产，缺乏移动端的内容触达与轻量化互动，无法形成PC与移动的双端流量闭环。")

    # 4. 竞品二：即梦AI (字节跳动)
    add_content_slide(prs, "竞品二：即梦AI (字节跳动)", [
        "PC端功能：一站式AI创作平台（智能画布、故事创作、运镜控制）。",
        "移动端有无：有移动端App（即梦AI）。",
        "短剧专属功能：无完整的短剧时间线功能。",
        "移动端功能：支持文生图、图生图、文/图生视频（单次生成视频片段），AI分身、社区作品“做同款”。",
        "PC与移动端关系：数据互通。移动端用于轻量级灵感生成与社区分享，PC端负责复杂的智能画布拼图与短剧分镜。导流至抖音体系获取流量扶持。"
    ], "即梦移动端重在“单镜头生成与社区灵感”，并未将“短剧生产全流”搬上手机，其闭环依赖字节抖音生态。")

    # 5. 竞品三：可灵AI (快手)
    add_content_slide(prs, "竞品三：可灵AI (快手)", [
        "PC端功能：视频生成大模型，支持多主体参考、智能分镜等影视级工业流。",
        "移动端有无：有移动端App（可灵AI）。",
        "短剧专属功能：缺乏长线短剧剧本解析与剪辑台。",
        "移动端功能：文生/图生视频、最长3分钟视频续写、运镜控制、一键同款。适合生成单片段的口型同步素材。",
        "PC与移动端关系：账号互通。移动端主打碎片化生成和素材收集，成品直接流向快手短视频生态进行变现。"
    ], "可灵AI移动端依然是“素材生产工具”，视频续写功能强大，但无法在移动端完成多集短剧的统筹与消费。")

    # 6. 竞品四：纳米 (360)
    add_content_slide(prs, "竞品四：纳米 (360)", [
        "PC端功能：纳米漫剧流水线（工业级AI漫剧生产平台，三维场景+时间线引擎）。",
        "移动端有无：有移动端App（纳米AI、纳米P视频）。",
        "短剧专属功能：移动端无工业级漫剧流水线入口。",
        "移动端功能：纳米AI主打超级智能体与搜索；纳米P视频主打一句话生视频、多图保角色、首尾帧转场。",
        "PC与移动端关系：PC端面向企业和专业承制方（B端），移动端面向C端提供P图、轻视频合成及AI搜索，双端定位割裂。"
    ], "纳米在PC端做重（企业流水线），在移动端做轻（搜索与短视频美化），并未将短剧专业创作下放移动端。")

    # 7. 竞品五：天工 (昆仑万维)
    add_content_slide(prs, "竞品五：天工 (昆仑万维)", [
        "PC端功能：天工短剧工作台（Agent驱动的自动化创作，一键生成分镜与成片）。",
        "移动端有无：有移动端App（天工超级智能体/天工AI助手）。",
        "短剧专属功能：无“短剧工作台”入口。",
        "移动端功能：主打文档、PPT、表格生成，AI搜索与对话，长文本阅读，语音合成。",
        "PC与移动端关系：移动端聚焦通用办公与助理场景，完全剥离了PC端特供的影视/短剧创作流。"
    ], "天工的短剧核心竞争力全部留在PC端，移动端彻底让位于通用办公与对话AI，无短剧创作生态。")

    # 8. 竞品六：万兴 (万兴科技)
    add_content_slide(prs, "竞品六：万兴剧厂 (万兴科技)", [
        "PC端功能：万兴剧厂（全链路Agent创作支持，剧本-分镜-剪辑全闭环）。",
        "移动端有无：万兴旗下有移动端App（万兴喵影），但“万兴剧厂”无独立移动端。",
        "短剧专属功能：移动端无。",
        "移动端功能：万兴喵影仅作为传统的视频剪辑工具，辅以基础AI包装特效。",
        "PC与移动端关系：万兴剧厂主攻Web/PC端的高效出片，与移动端传统剪辑软件未形成深度的AI短剧协同生态。"
    ], "万兴在短剧赛道主打PC端生产力工具赋能，移动端未跟进专用的短剧生成与消费平台。")

    # 9. 总结：PC端与移动端的关系映射
    add_content_slide(prs, "六大平台PC与移动端关系总结", [
        "1. 定位区隔：PC端无一例外是“重度生产力”（多镜头编排、角色管理、资产沉淀）；移动端则退化为“单点工具”（单镜生成、文本助手、搜索）。",
        "2. 流量引流：巨头（字节即梦、快手可灵）利用移动端App生成素材，随后直接引流至自家短视频App（抖音/快手）进行消费与流量扶持。",
        "3. 生态剥离：独立厂商（纳米、天工、万兴）在移动端放弃了短剧专业工作流，转向做大众化的通用AI助手或修图剪辑。",
        "4. 核心痛点：所有平台的移动端都缺乏“短剧”这个产品形态的连贯体验，用户无法在手机上完成“写剧本-选角色-生成多集-在线观看”的完整闭环。"
    ], "移动端当前是一片“单点素材生成”的红海，但却是“一站式短剧创作与消费”的蓝海。")

    # 10. 我们的移动端差异化竞争策略
    add_content_slide(prs, "我们的移动端差异化竞争策略", [
        "【策略1：前店后厂，消费与创作闭环】",
        "突破“纯工具”定位，移动端引入“类TikTok的短剧播放器”。用户在看剧时，可一键“改写结局”或“生成同款”，实现消费到创作的无缝转化。",
        "【策略2：卡牌化/对话式极简交互】",
        "放弃PC端复杂的时间线剪辑。在移动端采用“聊天式推进”或“卡牌翻页式”分镜管理，降低C端用户的短剧编排门槛。",
        "【策略3：云端资产库互通】",
        "PC端负责创建复杂的“专属AI演员和场景资产”，移动端直接调用云端资产进行快速的碎片化短剧续写与拼装。",
        "【策略4：社区共创与互动玩法】",
        "推出移动端专有的“接龙短剧”功能。官方发起第一集，用户用移动端轻量工具生成后续剧情，通过流量扶持与打赏分成实现拉新。"
    ], "避开与PC端拼“工业级生成能力”，在移动端主打“消费带动轻创作、卡牌化极简交互、社区剧情接龙”，形成降维打击。")

    prs.save('Mobile_Competitor_Analysis.pptx')
    print("PPT generated successfully!")

if __name__ == '__main__':
    main()
