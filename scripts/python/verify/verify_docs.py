# -*- coding: utf-8 -*-
from docx import Document
import os

files = [
    r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\短剧编剧Prompt智能增强系统_HermesAgent多轮对话版.docx',
    r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\视频生成Prompt智能增强系统_HermesAgent多轮对话版.docx'
]

for f in files:
    print(f'===== {os.path.basename(f)} =====')
    doc = Document(f)
    
    paras = [p for p in doc.paragraphs if p.text.strip()]
    print(f'总段落数（非空）: {len(paras)}')
    print(f'表格数: {len(doc.tables)}')
    
    img_count = 0
    for rel in doc.part.rels.values():
        if 'image' in rel.reltype:
            img_count += 1
    print(f'图片数: {img_count}')
    
    fonts = set()
    for p in doc.paragraphs[:80]:
        for run in p.runs:
            if run.font.name:
                fonts.add(run.font.name)
            if run._element.rPr is not None:
                rFonts = run._element.rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
                if rFonts is not None:
                    ea = rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia')
                    if ea:
                        fonts.add('eastAsia:' + ea)
    print(f'使用字体: {fonts}')
    
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
    print(f'标题总数: {len(headings)}')
    print('一级标题（章）:')
    for h in headings:
        if h.startswith('第') and '章' in h:
            print(f'  - {h}')
    
    full_text = ' '.join([p.text for p in paras])
    keywords = ['Hermes', '多轮对话', '上下文', 'Skills', '提示词', '帧对齐', '时间轴', '参考图', '负面提示词', '实际场景', '泳道', '架构', '分镜', '镜头', '剧本', '大纲']
    print('关键词覆盖:')
    for kw in keywords:
        cnt = full_text.count(kw)
        status = 'PASS' if cnt > 0 else 'MISS'
        print(f'  [{status}] {kw}: {cnt}次')
    
    # 检查提示词是否中文
    has_english_prompt = False
    for p in paras:
        if 'prompt' in p.text.lower() and len(p.text) > 50:
            # 检查是否大部分是英文
            english_chars = sum(1 for c in p.text if c.isascii() and c.isalpha())
            if english_chars > len(p.text) * 0.5:
                has_english_prompt = True
                break
    print(f'提示词中文展示: {"PASS" if not has_english_prompt else "CHECK"}')
    
    print()
