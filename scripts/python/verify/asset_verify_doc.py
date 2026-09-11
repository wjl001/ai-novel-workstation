# -*- coding: utf-8 -*-
"""
验证生成的Word文档
"""
from docx import Document
from docx.oxml.ns import qn
import os

DOC_PATH = r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\doc_asset_platform\AI内容生成平台_统一真人资产管理与多厂商适配方案.docx'

doc = Document(DOC_PATH)

print("=" * 60)
print("文档验证报告")
print("=" * 60)

# 1. 基本统计
para_count = len(doc.paragraphs)
table_count = len(doc.tables)
print(f"\n【基本统计】")
print(f"  段落总数: {para_count}")
print(f"  表格总数: {table_count}")

# 2. 图片统计
image_count = 0
for rel in doc.part.rels.values():
    if "image" in rel.reltype:
        image_count += 1
print(f"  图片总数: {image_count}")

# 3. 字体检查
print(f"\n【字体检查】")
font_issues = []
checked_runs = 0
yahei_runs = 0

for para in doc.paragraphs:
    for run in para.runs:
        checked_runs += 1
        # 检查中文字体
        rPr = run._element.find(qn('w:rPr'))
        if rPr is not None:
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is not None:
                east_asia = rFonts.get(qn('w:eastAsia'))
                ascii_font = rFonts.get(qn('w:ascii'))
                if east_asia and '雅黑' in east_asia:
                    yahei_runs += 1
                elif east_asia and '雅黑' not in east_asia:
                    font_issues.append(f"段落中文字体非雅黑: {east_asia} (文本: {run.text[:30]})")
            else:
                font_issues.append(f"段落未设置中文字体 (文本: {run.text[:30]})")

# 检查表格中的字体
table_font_issues = []
table_runs = 0
table_yahei = 0
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    table_runs += 1
                    rPr = run._element.find(qn('w:rPr'))
                    if rPr is not None:
                        rFonts = rPr.find(qn('w:rFonts'))
                        if rFonts is not None:
                            east_asia = rFonts.get(qn('w:eastAsia'))
                            if east_asia and '雅黑' in east_asia:
                                table_yahei += 1
                            elif east_asia and '雅黑' not in east_asia:
                                table_font_issues.append(f"表格字体非雅黑: {east_asia}")

print(f"  正文run总数: {checked_runs}")
print(f"  正文微软雅黑run数: {yahei_runs}")
print(f"  正文字体问题数: {len(font_issues)}")
print(f"  表格run总数: {table_runs}")
print(f"  表格微软雅黑run数: {table_yahei}")
print(f"  表格字体问题数: {len(table_font_issues)}")

if font_issues:
    print("\n  正文字体问题示例:")
    for issue in font_issues[:5]:
        print(f"    - {issue}")

# 4. 标题结构检查
print(f"\n【标题结构】")
headings = []
for para in doc.paragraphs:
    if para.style.name.startswith('Heading'):
        headings.append((para.style.name, para.text))

for style, text in headings:
    print(f"  {style}: {text}")

print(f"\n  标题总数: {len(headings)}")

# 5. 内容完整性检查
print(f"\n【内容完整性检查】")
full_text = '\n'.join([p.text for p in doc.paragraphs])
key_sections = [
    '文档概述', '背景与问题', '目标与范围', '术语定义',
    '现状分析', '功能点设计', '统一真人资产管理中心',
    '多厂商资产库自动同步', '生文流程', '生图流程', '生视频流程',
    '流程图与泳道图', '技术方案设计', '系统架构设计', '核心数据模型',
    '接口设计', '厂商适配器设计', '开发工作量评估', '任务拆解',
    '排期建议', '风险与应对', '总结与后续建议'
]

missing = []
for section in key_sections:
    if section not in full_text:
        missing.append(section)

if missing:
    print(f"  缺失的关键章节: {missing}")
else:
    print(f"  所有关键章节均已包含 ✓")

# 6. 表格内容检查
print(f"\n【表格内容检查】")
for i, table in enumerate(doc.tables):
    rows = len(table.rows)
    cols = len(table.columns)
    first_cell = table.rows[0].cells[0].text[:20]
    print(f"  表格{i+1}: {rows}行 x {cols}列, 首格: {first_cell}")

# 7. 文件大小
file_size = os.path.getsize(DOC_PATH)
print(f"\n【文件信息】")
print(f"  文件大小: {file_size / 1024:.1f} KB")
print(f"  文件路径: {DOC_PATH}")

print("\n" + "=" * 60)
print("验证完成")
print("=" * 60)
