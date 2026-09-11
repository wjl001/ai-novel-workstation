# -*- coding: utf-8 -*-
from docx import Document
from docx.oxml.ns import qn

doc = Document(r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\AI润色优化_固定模板对比测试方案_字数差异化版.docx')

check_indices = [4, 31, 32, 50, 51, 52]
for i in check_indices:
    p = doc.paragraphs[i]
    print(f'[{i}] style={p.style.name}, text={p.text[:40]}')
    for r in p.runs:
        rPr = r._element.find(qn('w:rPr'))
        font_info = ''
        if rPr is not None:
            rf = rPr.find(qn('w:rFonts'))
            if rf is not None:
                ea = rf.get(qn('w:eastAsia'))
                asc = rf.get(qn('w:ascii'))
                font_info += f' eastAsia={ea} ascii={asc}'
            sz = rPr.find(qn('w:sz'))
            if sz is not None:
                font_info += f' sz={sz.get(qn("w:val"))}'
            b = rPr.find(qn('w:b'))
            if b is not None:
                font_info += ' bold'
        print(f'    run: {r.text[:30]} |{font_info}')

print()
print('=== 表格5（优化后提示词）字体 ===')
t = doc.tables[5]
for p in t.rows[0].cells[0].paragraphs[:3]:
    for r in p.runs:
        rPr = r._element.find(qn('w:rPr'))
        font_info = ''
        if rPr is not None:
            rf = rPr.find(qn('w:rFonts'))
            if rf is not None:
                font_info += f' eastAsia={rf.get(qn("w:eastAsia"))}'
            sz = rPr.find(qn('w:sz'))
            if sz is not None:
                font_info += f' sz={sz.get(qn("w:val"))}'
        print(f'  {r.text[:40]} |{font_info}')
