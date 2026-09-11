from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def set_font(run, font_name='微软雅黑', size=12, bold=False, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_heading(doc, text, level, font_name='微软雅黑'):
    h = doc.add_heading('', level=level)
    run = h.add_run(text)
    size = 18 if level == 0 else (16 if level == 1 else (14 if level == 2 else 12))
    set_font(run, font_name, size=size, bold=True)
    return h

def add_paragraph(doc, text, font_name='微软雅黑', size=11, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, font_name, size=size, bold=bold)
    return p

doc = Document()

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('AI短剧 v1.0.2 版本需求文档')
set_font(run, size=22, bold=True)

# 文档说明
add_heading(doc, '文档说明', level=1)
add_paragraph(doc, '本文档为AI短剧创作平台v1.0.2迭代版本功能需求说明书。本次迭代旨在进一步优化用户创作体验，提升主体设定的灵活性与专业度，强化模型平台的合规性与多样性，并完善会员中心的商业化闭环。')
add_paragraph(doc, '版本信息：V1.0.2', bold=True)
add_paragraph(doc, '适配终端：PC端Web平台', bold=True)
add_paragraph(doc, '迭代目标：主体设定深度优化、模型平台能力扩展、分镜视频生成体验升级、会员中心明细完善、首页引导与计费策略优化。', bold=True)

# 1. 底层服务 (Briefly mention or keep v1.0.0 ones if relevant, but here we focus on v1.0.2 changes)
add_heading(doc, '1. 底层服务', level=1)
add_paragraph(doc, '（沿用 v1.0.0 异步自动重连机制，保障高并发下的连接稳定性）')

# 2. 灵感创作与首页
add_heading(doc, '2. 灵感创作与首页', level=1)
add_heading(doc, '2.1 灵感生成页面视频比例优化', level=2)
add_paragraph(doc, '功能详细描述：在灵感生成页面新增 9:16（竖屏）和 16:9（横屏）比例选择，适配主流短视频与剧集平台。')
add_heading(doc, '2.2 首页热门题材计费策略', level=2)
add_paragraph(doc, '功能详细描述：热门题材支持平台损耗模式，用户体验不扣除个人算力豆，由系统后台统一配置。')

# 3. 会员中心
add_heading(doc, '3. 会员中心', level=1)
add_heading(doc, '3.1 充值记录查询', level=2)
add_paragraph(doc, '功能详细描述：新增充值记录模块，展示时间、金额、套餐及支付状态。')

# 4. 主体设置
add_heading(doc, '4. 主体设置', level=1)
add_heading(doc, '4.1 主体新增与删除', level=2)
add_paragraph(doc, '功能详细描述：支持用户手动在库中新增或删除主体资产。')
add_heading(doc, '4.2 角色音频参考', level=2)
add_paragraph(doc, '功能详细描述：新增上传音频参考文件（mp3/wav）功能，为AI提供声画统一基准。')
add_heading(doc, '4.3 生成分辨率配置', level=2)
add_paragraph(doc, '功能详细描述：支持预设主体生成的分辨率规格（512/1024等）。')
add_heading(doc, '4.4 描述词非必填逻辑', level=2)
add_paragraph(doc, '功能详细描述：上传站外图片后，文字描述转为可选，提升创建效率。')
add_heading(doc, '4.5 异步生成机制', level=2)
add_paragraph(doc, '功能详细描述：生成主体图片改为后台异步处理，支持弹窗关闭后继续执行。')
add_heading(doc, '4.6 场景同步策略', level=2)
add_paragraph(doc, '功能详细描述：场景类主体默认勾选同步至视频资产库。')

# 5. 剧本创作
add_heading(doc, '5. 剧本创作', level=1)
add_heading(doc, '5.1 生成进度时间预估', level=2)
add_paragraph(doc, '功能详细描述：大纲生成阶段实时展示 AI 构思的预计剩余时间。')

# 6. 分镜视频
add_heading(doc, '6. 分镜视频', level=1)
add_heading(doc, '6.1 失败原因反馈', level=2)
add_paragraph(doc, '功能详细描述：批量生成失败时显式展示错误原因（敏感词/超时/余额不足等）。')
add_heading(doc, '6.2 分辨率记录记忆', level=2)
add_paragraph(doc, '功能详细描述：自动记忆并默认填充用户上一次成功选择的分辨率记录。')
add_heading(doc, '6.3 算力豆消耗归集', level=2)
add_paragraph(doc, '功能详细描述：将文本转换与视频渲染算力豆合并展示，优化计费感知。')
add_heading(doc, '6.4 布局与全集进度优化', level=2)
add_paragraph(doc, '功能详细描述：分镜配置布局调整，合成全集仅保留统一进度百分比。')

# 7. 模型平台
add_heading(doc, '7. 模型平台', level=1)
add_heading(doc, '7.1 敏感词拦截（一期）', level=2)
add_paragraph(doc, '功能详细描述：实现输入与输出的双向合规拦截，建立风控日志。')
add_heading(doc, '7.2 豆包 Seedance 2.0 接入', level=2)
add_paragraph(doc, '功能详细描述：新增 doubao-seedance-2-0-fast 与 doubao-seedance-2-0-mini 模型支持。')

# 9. 版本迭代总结
add_heading(doc, '9. 版本迭代总结', level=1)
add_paragraph(doc, '本次 v1.0.2 版本通过主体精细化管控、模型扩展与交互反馈优化，进一步巩固了 AI 短剧创作的工业化底座。')

doc.save(r'd:\phpstudy_pro\WWW\ai-novel-workstation1.0.1\AI短剧 v1.0.2版本需求文档.docx')
print("v1.0.2 PRD generated successfully.")
