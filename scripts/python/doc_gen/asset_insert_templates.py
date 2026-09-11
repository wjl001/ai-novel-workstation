# -*- coding: utf-8 -*-
"""
在AI润色优化文档中插入固定提示词模板章节
插入位置："二、测试方法与三大模型介绍"之前
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import copy
import shutil

SRC = r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\AI润色优化_固定模板对比测试方案_字数差异化版.docx'
DST = r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\AI润色优化_固定模板对比测试方案_字数差异化版_含固定模板.docx'

# 先备份
shutil.copy2(SRC, SRC + '.bak')
print('已备份原文件')

doc = Document(SRC)

FONT = '微软雅黑'

def set_run_font(run, size_pt=11, bold=False, color=None):
    """设置run字体，size_pt单位是磅，color可以是RGBColor或十六进制字符串"""
    run.font.name = FONT
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        if isinstance(color, str):
            color = RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))
        run.font.color.rgb = color
    r = run._element
    rPr = r.find(qn('w:rPr'))
    if rPr is None:
        rPr = parse_xml(f'<w:rPr {nsdecls("w")}></w:rPr>')
        r.insert(0, rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="{FONT}" w:ascii="{FONT}" w:hAnsi="{FONT}"/>')
        rPr.insert(0, rFonts)
    else:
        rFonts.set(qn('w:eastAsia'), FONT)
        rFonts.set(qn('w:ascii'), FONT)
        rFonts.set(qn('w:hAnsi'), FONT)

def make_paragraph(text, size_pt=11, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT,
                    space_before=0, space_after=0, color=None):
    """创建一个新段落"""
    p = parse_xml(f'<w:p {nsdecls("w")}></w:p>')
    # 段落属性
    pPr = parse_xml(f'<w:pPr {nsdecls("w")}></w:pPr>')
    if align == WD_ALIGN_PARAGRAPH.CENTER:
        jc = parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>')
        pPr.append(jc)
    spacing = parse_xml(f'<w:spacing {nsdecls("w")} w:before="{space_before*20}" w:after="{space_after*20}" w:line="360" w:lineRule="auto"/>')
    pPr.append(spacing)
    p.append(pPr)
    # run
    r = parse_xml(f'<w:r {nsdecls("w")}></w:r>')
    rPr = parse_xml(f'<w:rPr {nsdecls("w")}></w:rPr>')
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="{FONT}" w:ascii="{FONT}" w:hAnsi="{FONT}"/>')
    rPr.append(rFonts)
    sz = parse_xml(f'<w:sz {nsdecls("w")} w:val="{size_pt*2}"/>')
    rPr.append(sz)
    szCs = parse_xml(f'<w:szCs {nsdecls("w")} w:val="{size_pt*2}"/>')
    rPr.append(szCs)
    if bold:
        b = parse_xml(f'<w:b {nsdecls("w")}/>')
        rPr.append(b)
        bCs = parse_xml(f'<w:bCs {nsdecls("w")}/>')
        rPr.append(bCs)
    if color:
        color_elem = parse_xml(f'<w:color {nsdecls("w")} w:val="{color}"/>')
        rPr.append(color_elem)
    r.append(rPr)
    t = parse_xml(f'<w:t {nsdecls("w")} xml:space="preserve">{text}</w:t>')
    r.append(t)
    p.append(r)
    return p

def make_table_with_content(headers, rows, col_widths_cm=None):
    """创建一个带内容的表格，返回tbl元素"""
    # 用python-docx创建临时表格再复制XML
    tmp_doc = Document()
    table = tmp_doc.add_table(rows=1+len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, size_pt=10, bold=True)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        # 表头背景色
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="D9E2F3"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    # 数据行
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx+1].cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(val))
            set_run_font(run, size_pt=9.5)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    if col_widths_cm:
        for i, w in enumerate(col_widths_cm):
            for row in table.rows:
                row.cells[i].width = Cm(w)

    return table._tbl

def make_prompt_table(title, content_lines):
    """创建提示词模板表格（1列，带标题行）"""
    tmp_doc = Document()
    table = tmp_doc.add_table(rows=2, cols=1)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 标题行
    cell0 = table.rows[0].cells[0]
    cell0.text = ''
    p0 = cell0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run0 = p0.add_run(title)
    set_run_font(run0, size_pt=11, bold=True, color='1F4E79')
    cell0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    shading0 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="D6E4F0"/>')
    cell0._tc.get_or_add_tcPr().append(shading0)

    # 内容行
    cell1 = table.rows[1].cells[0]
    cell1.text = ''
    for i, line in enumerate(content_lines):
        if i == 0:
            p = cell1.paragraphs[0]
        else:
            p = cell1.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.space_before = Pt(2)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.3
        # 占位符行用蓝色加粗
        if line.startswith('【') and '】' in line[:6]:
            run = p.add_run(line)
            set_run_font(run, size_pt=9.5, bold=True, color='2E75B6')
        else:
            run = p.add_run(line)
            set_run_font(run, size_pt=9.5)
    cell1.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    return table._tbl


# ============================================================
# 找到插入位置："二、测试方法与三大模型介绍"
# ============================================================
target_para = None
for p in doc.paragraphs:
    if '二、测试方法' in p.text:
        target_para = p
        break

if target_para is None:
    print('ERROR: 未找到插入位置')
    exit(1)

print(f'找到插入位置: {target_para.text}')

# ============================================================
# 构建要插入的内容
# ============================================================
elements_to_insert = []

# 大标题
elements_to_insert.append(make_paragraph(
    '【补充】AI润色优化固定提示词模板（点击按钮后核心处理逻辑）',
    size_pt=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
    space_before=12, space_after=8, color='1F4E79'
))

# 说明文字
elements_to_insert.append(make_paragraph(
    '用户点击"AI润色优化"按钮后，系统按以下逻辑处理：①自动识别用户输入的主体类型（角色/场景/道具）；②调用对应类别的固定优化模板；③将用户描述中的关键词提取并填入模板占位符；④输出完整的专业级提示词。以下为三大类别的固定优化提示词模板，模板中【】标注的为动态填充占位符。',
    size_pt=11, space_after=6
))

# 调用逻辑表
elements_to_insert.append(make_paragraph(
    '表补-1  AI润色按钮点击后处理逻辑',
    size_pt=11, bold=True, space_before=6, space_after=4
))
logic_table = make_table_with_content(
    ['步骤', '处理动作', '说明', '耗时'],
    [
        ['① 主体识别', '关键词匹配+规则判断', '识别用户输入属于角色/场景/道具哪一类，四视图角色单独识别', '<0.1秒'],
        ['② 模板选择', '调用对应类别固定模板', '角色四视图模板 / 场景模板 / 道具模板', '<0.1秒'],
        ['③ 占位符填充', '提取用户描述填入模板', '将用户输入的核心描述填入【主体描述】【风格】等占位符', '<0.1秒'],
        ['④ 质量增强', '自动补充画质/光影/负面词', '统一追加8K超高清、极致写实、专业布光、负面提示词等标准段', '<0.1秒'],
        ['⑤ 输出结果', '返回完整优化后提示词', '前端展示优化后提示词，用户可直接用于生图或继续编辑', '<0.1秒'],
    ],
    col_widths_cm=[2.5, 3.5, 7, 2]
)
elements_to_insert.append(logic_table)
elements_to_insert.append(make_paragraph('', size_pt=6, space_after=4))

# ===== 模板一：角色四视图 =====
elements_to_insert.append(make_paragraph(
    '模板一：角色四视图优化提示词模板',
    size_pt=14, bold=True, space_before=10, space_after=6, color='1F4E79'
))
elements_to_insert.append(make_paragraph(
    '适用场景：用户输入角色描述且需要四视图（正面/侧面/背面/面部特写）时调用。点击AI润色按钮后，系统自动将用户描述填入以下模板。',
    size_pt=10.5, space_after=4
))

char_prompt_lines = [
    '【画质前缀】8K超高清，真实电影级【风格】人物摄影，极致写实，明暗过渡细腻，皮肤肌理真实，毛发丝缕质感，面料纹路清晰，纯白色干净背景，四视图设定表格式，从左到右依次为面部特写、正面全身、侧面全身、背面全身，四角人物特征完全一致，人物气质【气质描述】。',
    '',
    '【主体描述】【年龄性别】，身姿【身姿描述】；五官【五官描述】，【面部细节】。',
    '',
    '【发型描述】【发型】，【发色发质】，【发饰细节】，发丝根根分明，四角一致。',
    '',
    '【服饰描述】【服装风格】，整体【色系】；【上衣描述】；【腰部配饰】；【手臂细节】；【下装描述】；【鞋履描述】，四角一致。',
    '',
    '①面部特写：肩部以上，正面朝向，双眼视线自然聚焦瞳孔对称，五官比例协调，重点展示面部细节。',
    '②正面全身：标准站立双臂下垂双脚并拢，全身完整无截断，面部与特写一致。',
    '③侧面全身：左侧面，轮廓清晰鼻梁挺直，高矮体型与正面一致。',
    '④背面全身：背面站立不回头，后脑和服饰背面完整，高矮与正面一致。',
    '',
    '负面提示词：人脸崩坏，五官扭曲，脸型变形，斗鸡眼斜视，瞳孔偏移，双眼不一致，多脸残脸，浓妆假面，锥子脸，过度磨皮，假白肤色，肤色不均色块，彩色怪异发色，发型发色不一致，塑料面料，粗糙布料，纹饰杂乱艳俗，服饰颜色款式不一致，现代服饰暴露服装，背景杂乱，多余人物道具，四角排列错误，缺少某视角，人物截断，比例失调，肢体畸形，卡通动漫，CG感过重，二次元，手绘感，模糊失真，低质量像素化，锯齿，色彩异常。',
]
elements_to_insert.append(make_prompt_table('角色四视图固定优化模板（点击AI润色后自动填充）', char_prompt_lines))
elements_to_insert.append(make_paragraph('', size_pt=6, space_after=4))

# 模板占位符说明表
elements_to_insert.append(make_paragraph(
    '表补-2  角色四视图模板占位符说明',
    size_pt=11, bold=True, space_before=4, space_after=4
))
char_placeholder_table = make_table_with_content(
    ['占位符', '填充来源', '示例值', '是否必填'],
    [
        ['【风格】', '用户输入+规则匹配', '古风 / 科幻 / 现代职场 / 历史', '是'],
        ['【气质描述】', '用户输入提取', '清冷灵动 / 威严庄重 / 干练精英', '是'],
        ['【年龄性别】', '用户输入提取', '25岁女性 / 40岁男性', '否（有默认）'],
        ['【身姿描述】', '规则生成', '利落英挺 / 身姿挺拔 / 伟岸', '否（有默认）'],
        ['【五官描述】', '用户输入提取+规则增强', '清纯靓丽 / 剑眉星目 / 精致秀气', '是'],
        ['【发型】', '用户输入提取', '高马尾 / 束发 / 短发 / 长发披肩', '是'],
        ['【服装风格】', '用户输入提取', '侠客服饰 / 西装 / 龙袍 / 机甲', '是'],
        ['【色系】', '规则匹配', '青色系 / 深灰 / 明黄 / 银黑', '否（有默认）'],
    ],
    col_widths_cm=[3, 3.5, 5.5, 2.5]
)
elements_to_insert.append(char_placeholder_table)
elements_to_insert.append(make_paragraph('', size_pt=6, space_after=4))

# ===== 模板二：场景 =====
elements_to_insert.append(make_paragraph(
    '模板二：场景优化提示词模板',
    size_pt=14, bold=True, space_before=10, space_after=6, color='1F4E79'
))
elements_to_insert.append(make_paragraph(
    '适用场景：用户输入场景描述（室内/室外/自然/城市等）时调用。系统自动按"前景-中景-背景"专业构图方式填充模板。',
    size_pt=10.5, space_after=4
))

scene_prompt_lines = [
    '【画质前缀】8K超高清，电影级场景摄影，极致写实，明暗层次丰富，真实材质质感，空间透视准确，氛围沉浸感强。',
    '',
    '【场景整体描述】【场景类型】场景，【建筑结构描述】，【空间特征描述】。',
    '',
    '前景：【前景物品描述】，【前景细节】，【前景光影效果】。',
    '中景：【中景区域描述】，【中景物品/人物】，【中景活动细节】。',
    '背景：【背景区域描述】，【背景元素】，【背景纵深细节】。',
    '',
    '光线：主光源来自【光源方向】，辅以【辅助光源】，形成明暗对比，角落自然暗角，营造【氛围描述】。',
    '色彩：【主色调】为主色调，搭配【点缀色】点缀，整体色调【色调评价】。',
    '构图：【构图方式】，视线引导至【视觉焦点】，前景中景背景层次分明，空间感强烈。',
    '',
    '负面提示词：透视错误，空间扭曲，比例失调，现代物品，塑料质感，材质虚假，光线平淡无层次，色彩杂乱，过曝过暗，模糊失真，低质量，卡通动漫，CG感过重，二次元，手绘感。',
]
elements_to_insert.append(make_prompt_table('场景固定优化模板（点击AI润色后自动填充）', scene_prompt_lines))
elements_to_insert.append(make_paragraph('', size_pt=6, space_after=4))

# 场景占位符说明表
elements_to_insert.append(make_paragraph(
    '表补-3  场景模板占位符说明',
    size_pt=11, bold=True, space_before=4, space_after=4
))
scene_placeholder_table = make_table_with_content(
    ['占位符', '填充来源', '示例值', '是否必填'],
    [
        ['【场景类型】', '用户输入提取', '古代客栈室内 / 现代都市夜景 / 未来科幻城市', '是'],
        ['【建筑结构描述】', '规则生成+用户提取', '木质结构榫卯梁架 / 摩天大楼玻璃幕墙', '是'],
        ['【前景物品描述】', '规则生成（按场景类型）', '方桌酒壶油灯 / 护栏车流光轨', '否（有默认）'],
        ['【中景区域描述】', '规则生成', '柜台货架 / 街道商铺 / 广场人群', '否（有默认）'],
        ['【背景区域描述】', '规则生成', '楼梯字画灯笼 / 远处天际线 / 巨型建筑', '否（有默认）'],
        ['【光源方向】', '规则匹配', '窗户侧光 / 霓虹灯光 / 顶部自然光', '否（有默认）'],
        ['【氛围描述】', '用户输入提取', '温馨热闹 / 繁华都市 / 赛博朋克未来', '是'],
        ['【主色调】', '规则匹配', '暖木色 / 冷蓝霓虹 / 金红宫廷', '否（有默认）'],
        ['【构图方式】', '规则生成', '一点透视 / 低角度仰拍 / 鸟瞰俯视', '否（有默认）'],
    ],
    col_widths_cm=[3, 3.5, 5.5, 2.5]
)
elements_to_insert.append(scene_placeholder_table)
elements_to_insert.append(make_paragraph('', size_pt=6, space_after=4))

# ===== 模板三：道具 =====
elements_to_insert.append(make_paragraph(
    '模板三：道具优化提示词模板',
    size_pt=14, bold=True, space_before=10, space_after=6, color='1F4E79'
))
elements_to_insert.append(make_paragraph(
    '适用场景：用户输入道具/物品描述（武器、法器、工具、饰品等）时调用。系统自动按"整体-材质-光影-构图"专业产品摄影方式填充模板。',
    size_pt=10.5, space_after=4
))

prop_prompt_lines = [
    '【画质前缀】8K超高清，产品级静物摄影，极致写实质感，专业布光，细节锐利清晰，【材质类型】质感真实，【氛围类型】氛围。',
    '',
    '【道具整体描述】【道具名称】，【整体形态描述】，【关键部件描述】，【核心特征描述】。',
    '',
    '材质表现：【主体材质描述】，【材质反光/纹理细节】，【使用痕迹/年代感】；【部件1材质】，【部件1细节】；【部件2材质】，【部件2细节】；【部件3材质】，【部件3细节】。',
    '',
    '光影：主光【主光角度】，塑造立体感，辅光柔化阴影，边缘光勾勒轮廓，背景【背景描述】，突出主体。',
    '构图：【摆放方式】，【占据画面位置】，周围留白充足，产品级展示效果。',
    '色彩：【主色】为主，【辅助色1】，【辅助色2】，整体色调【色调评价】。',
    '',
    '负面提示词：材质虚假，塑料感，金属反光错误，细节模糊，纹理不清，光线平淡，构图杂乱，背景干扰，现代元素，卡通动漫，CG感过重，低质量，像素化，【道具特有负面词】。',
]
elements_to_insert.append(make_prompt_table('道具固定优化模板（点击AI润色后自动填充）', prop_prompt_lines))
elements_to_insert.append(make_paragraph('', size_pt=6, space_after=4))

# 道具占位符说明表
elements_to_insert.append(make_paragraph(
    '表补-4  道具模板占位符说明',
    size_pt=11, bold=True, space_before=4, space_after=4
))
prop_placeholder_table = make_table_with_content(
    ['占位符', '填充来源', '示例值', '是否必填'],
    [
        ['【材质类型】', '规则匹配', '金属 / 水晶 / 木材 / 科技材质', '是'],
        ['【氛围类型】', '用户输入提取', '古董文物 / 魔幻神秘 / 未来科技', '是'],
        ['【道具名称】', '用户输入提取', '古代宝剑 / 魔法水晶球 / 科技能量手枪', '是'],
        ['【整体形态描述】', '用户输入提取', '剑身修长笔直 / 球体通透 / 流线型设计', '是'],
        ['【关键部件描述】', '规则生成（按道具类型）', '剑格剑柄剑首 / 底座 / 握把能量灯', '否（有默认）'],
        ['【主体材质描述】', '规则匹配', '精钢锻造 / 高透玻璃 / 哑光钛合金', '否（有默认）'],
        ['【主光角度】', '规则生成', '45度侧上方 / 正面柔光 / 顶部逆光', '否（有默认）'],
        ['【背景描述】', '规则生成', '渐变深灰到黑色 / 深黑纯色 / 虚化场景', '否（有默认）'],
        ['【摆放方式】', '规则生成', '斜向45度摆放 / 正面居中 / 悬浮展示', '否（有默认）'],
    ],
    col_widths_cm=[3, 3.5, 5.5, 2.5]
)
elements_to_insert.append(prop_placeholder_table)
elements_to_insert.append(make_paragraph('', size_pt=6, space_after=4))

# 模板对比总结表
elements_to_insert.append(make_paragraph(
    '表补-5  三大模板对比与适用场景',
    size_pt=11, bold=True, space_before=6, space_after=4
))
summary_table = make_table_with_content(
    ['模板类型', '识别关键词', '核心结构', '优化后字数', '适用模型'],
    [
        ['角色四视图', '角色、人物、四视图、正面侧面背面', '画质+人物+发型+服饰+四视角+负面词', '约400-500字', '所有生图模型'],
        ['场景', '场景、室内、室外、城市、风景、房间', '画质+整体+前中后景+光线+色彩+构图+负面词', '约300-400字', '所有生图模型'],
        ['道具', '道具、武器、物品、剑、球、枪、饰品', '画质+整体+材质+光影+构图+色彩+负面词', '约250-350字', '所有生图模型'],
    ],
    col_widths_cm=[2.5, 4, 5, 2, 2.5]
)
elements_to_insert.append(summary_table)
elements_to_insert.append(make_paragraph('', size_pt=8, space_after=8))

# 分隔线
elements_to_insert.append(make_paragraph(
    '———————————————————— 以下为原测试方案内容 ————————————————————',
    size_pt=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
    space_before=8, space_after=8, color='808080'
))

# ============================================================
# 执行插入：在目标段落之前依次插入所有元素
# ============================================================
target_element = target_para._element
for elem in elements_to_insert:
    target_element.addprevious(elem)

print(f'已插入 {len(elements_to_insert)} 个元素')

# 保存
doc.save(DST)
print(f'已保存: {DST}')
