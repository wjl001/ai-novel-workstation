#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""完整版V2 - 详细实施手册"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def sf(run, name='微软雅黑', size=11, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)

def H(doc, text, level=1):
    p = doc.add_paragraph()
    r = p.add_run(text)
    sf(r, size={1:18,2:15,3:13,4:12}.get(level,12), bold=True, color=(30,58,138))
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)

def P(doc, text, size=11, bold=False, indent=True):
    p = doc.add_paragraph()
    if indent: p.paragraph_format.first_line_indent = Pt(22)
    r = p.add_run(text)
    sf(r, size=size, bold=bold)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)

def B(doc, text, size=11, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(22 + level*22)
    r = p.add_run('● ' + text)
    sf(r, size=size)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(2)

def T(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i,h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text=''
        p = c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h); sf(r, size=10, bold=True, color=(255,255,255))
        s = OxmlElement('w:shd'); s.set(qn('w:fill'),'1E3A8A'); c._tc.get_or_add_tcPr().append(s)
    for ri,row in enumerate(rows):
        for ci,val in enumerate(row):
            c = t.rows[ri+1].cells[ci]; c.text=''
            p = c.paragraphs[0]; r = p.add_run(str(val)); sf(r, size=10)
            if ri%2==1:
                s = OxmlElement('w:shd'); s.set(qn('w:fill'),'F0F4FF'); c._tc.get_or_add_tcPr().append(s)
    if widths:
        for i,w in enumerate(widths):
            for row in t.rows: row.cells[i].width = Cm(w)

def FB(doc, title, content='', color=(30,58,138)):
    t = doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; c.text=''
    p = c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); sf(r, size=11, bold=True, color=(255,255,255))
    s = OxmlElement('w:shd'); s.set(qn('w:fill'),'%02X%02X%02X'%color); c._tc.get_or_add_tcPr().append(s)
    if content:
        p2 = c.add_paragraph(); p2.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(content); sf(r2, size=9, color=(255,255,255))

def AR(doc):
    p = doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('↓'); sf(r, size=16, bold=True, color=(30,58,138))

def CB(doc, text, title=''):
    if title:
        p = doc.add_paragraph(); r = p.add_run(title); sf(r, size=10, bold=True, color=(100,100,100))
    t = doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; c.text=''
    s = OxmlElement('w:shd'); s.set(qn('w:fill'),'F8F9FA'); c._tc.get_or_add_tcPr().append(s)
    for line in text.split('\n'):
        p = c.add_paragraph(); r = p.add_run(line); sf(r, size=10, color=(50,50,50))
        p.paragraph_format.space_after = Pt(2)

doc = Document()
style = doc.styles['Normal']
style.font.name = '微软雅黑'; style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

# 封面
for _ in range(4): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('人工智能短剧创作系统'); sf(r,size=28,bold=True,color=(30,58,138))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('图片生成提示词优化产品方案'); sf(r,size=24,bold=True,color=(30,58,138))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('—— 基于 Hermes Agent 多轮对话智能优化 ——'); sf(r,size=16,color=(100,100,100))
for _ in range(2): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('【完整版 · 详细实施手册】'); sf(r,size=18,bold=True,color=(220,38,38))
for _ in range(5): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('面向：产品部 · 研发部 · 测试部 · 运营部'); sf(r,size=12,color=(100,100,100))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('版本：V2.0（Hermes Agent 多轮对话详细版）'); sf(r,size=12,color=(100,100,100))
doc.add_page_break()

# 目录
H(doc,'目录',1)
for item in ['一、项目背景与问题分析','二、产品目标与核心价值','三、产品架构设计（含详细架构图）',
'四、Hermes Agent 多轮对话优化机制（详细实施）','五、提示词 Skills 模板库详解（92套完整列举）',
'六、角色三视图设定专项设计（详细实施）','七、主体设置页面交互设计（含用户确认）',
'八、业务流程设计（含详细流程图）','九、跨角色协作泳道设计','十、技术实现方案',
'十一、测试验收方案','十二、项目排期与风险管理']:
    p=doc.add_paragraph(); r=p.add_run(item); sf(r,size=12); p.paragraph_format.space_after=Pt(6)
doc.add_page_break()

# 第一章
H(doc,'一、项目背景与问题分析',1)
H(doc,'1.1 项目背景',2)
P(doc,'人工智能短剧创作系统是一款面向短剧创作者的全流程人工智能辅助工具，覆盖剧本生成、大纲拆解、主体资产管理、分镜生成、视频合成等核心环节。其中，主体资产管理模块（角色、场景、道具）是整个系统的视觉基础，其图片生成质量直接决定了最终短剧视频的视觉效果和用户体验。')
P(doc,'当前系统已实现主体图片的批量生成与单张重生成功能，但在实际使用过程中，用户反馈图片生成质量不稳定，尤其在人物角色图片方面存在多种典型问题，严重影响创作效率和成品质量。')
P(doc,'为解决上述问题，本方案提出基于 Hermes Agent 底层多轮对话的提示词智能优化方案。与传统基于规则的模板匹配不同，本方案借鉴短剧编剧模块（骨架生成器、大纲填充器、主模板）和视频生成模块（分镜拆解、镜头拆解）的成熟多轮对话经验，通过 Hermes Agent 底层进行多轮隐式对话，反复迭代优化提示词。')
P(doc,'核心特点：界面不显示对话内容，所有优化过程在 Hermes 底层通过多轮隐式对话完成，用户仅需确认最终优化结果。同时，多轮对话的上下文会自动融入当前剧集和剧本的相关信息，确保优化方向与剧情需求一致。')

H(doc,'1.2 现存问题详细分析',2)
P(doc,'通过对近一个月用户生成图片的抽样分析与用户反馈收集，我们梳理出以下四类高频问题：')
T(doc,['问题类型','表现描述','发生频率','影响程度','典型案例'],
[['图片生成精度不高','画面模糊、细节丢失、纹理不清晰、边缘锯齿明显','约45%','高','角色服装纹理模糊，场景建筑细节丢失'],
['人物图片色彩异常','皮肤出现非自然色块、服饰颜色溢出、背景色污染主体','约32%','高','角色脸部出现紫色色块，红色衣服染到背景'],
['斗鸡眼问题','双眼视线不聚焦、瞳孔位置偏移、双眼朝向不一致','约28%','极高','角色双眼看向不同方向，瞳孔向内靠拢'],
['人脸崩坏','面部五官扭曲、脸型变形、五官比例失调、多脸或残脸','约25%','极高','角色嘴巴歪斜，眼睛一大一小，出现两张脸']],
widths=[2.5,4.5,1.8,1.8,3.5])

H(doc,'1.3 问题根因总结',2)
B(doc,'第一点：提示词结构过于简单。现有提示词仅包含基础描述加通用风格词，缺少结构化的质量控制层、特征锁定层、负面排除层。')
B(doc,'第二点：缺少迭代优化机制。当前提示词为一次性生成，缺少多轮迭代优化过程，无法根据问题特征进行针对性调整。')
B(doc,'第三点：缺少剧情上下文。提示词优化时未考虑主体所属的剧集、剧本、场景情绪等上下文信息，导致优化方向与剧情需求脱节。')
B(doc,'第四点：无质量检测与自动优化闭环。生成后缺少自动质量检测环节，问题图片无法被自动识别和重试优化。')
doc.add_page_break()

# 第二章
H(doc,'二、产品目标与核心价值',1)
H(doc,'2.1 产品目标',2)
T(doc,['目标维度','具体指标','当前基线','目标值'],
[['图片一次成功率','首次生成即满足用户需求的比例','约55%','≥85%'],
['人脸崩坏率','人物图片出现面部五官扭曲的比例','约25%','≤5%'],
['斗鸡眼发生率','人物图片出现双眼视线异常的比例','约28%','≤3%'],
['色彩异常率','图片出现非自然色块或颜色溢出的比例','约32%','≤8%'],
['用户重生成次数','平均每个主体需要重生成的次数','2.8次','≤1.2次'],
['提示词优化耗时','系统完成提示词优化的平均时间','无优化','≤8秒']],
widths=[3,5,2.5,2.5])
H(doc,'2.2 核心价值',2)
B(doc,'用户价值：大幅提升图片生成质量与一次成功率，减少反复调整和重生成的时间成本。')
B(doc,'产品价值：建立"Hermes Agent 多轮对话提示词优化引擎加质量检测闭环"的技术壁垒。')
B(doc,'团队价值：通过隐式优化机制和用户确认机制，降低生成结果的不可控性。')
B(doc,'技术价值：探索 Hermes Agent 在视觉生成提示词优化领域的应用，为后续场景积累经验。')
doc.add_page_break()

print('前两章完成，继续第三章...')
doc.save(r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\temp_doc.docx')
print('临时保存成功')
