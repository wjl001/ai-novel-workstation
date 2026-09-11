#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成基于 Hermes Agent 多轮对话的图片生成提示词优化产品方案 - 完整版"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_font(run, name='微软雅黑', size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    sizes = {1: 18, 2: 15, 3: 13, 4: 12}
    set_font(run, size=sizes.get(level, 12), bold=True, color=(30, 58, 138))
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_para(doc, text, size=11, bold=False, indent=True, color=None):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(22)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    return p

def add_bullet(doc, text, size=11, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(22 + level * 22)
    run = p.add_run('● ' + text)
    set_font(run, size=size)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(2)
    return p

def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    # 表头
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_font(run, size=10, bold=True, color=(255, 255, 255))
        # 设置背景色
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), '1E3A8A')
        cell._tc.get_or_add_tcPr().append(shading)
    # 数据行
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_font(run, size=10)
            if r_idx % 2 == 1:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'F0F4FF')
                cell._tc.get_or_add_tcPr().append(shading)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    return table

def add_flow_box(doc, title, content, color=(30, 58, 138)):
    """添加流程图框（用表格模拟）"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_font(run, size=11, bold=True, color=(255, 255, 255))
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), '%02X%02X%02X' % color)
    cell._tc.get_or_add_tcPr().append(shading)
    if content:
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = p2.add_run(content)
        set_font(run2, size=9, color=(255, 255, 255))
    return table

def add_arrow(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('↓')
    set_font(run, size=16, bold=True, color=(30, 58, 138))

# 创建文档
doc = Document()

# 设置默认字体
style = doc.styles['Normal']
style.font.name = '微软雅黑'
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

# ========== 封面 ==========
for _ in range(4):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('人工智能短剧创作系统')
set_font(run, size=28, bold=True, color=(30, 58, 138))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('图片生成提示词优化产品方案')
set_font(run, size=24, bold=True, color=(30, 58, 138))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('—— 基于 Hermes Agent 多轮对话智能优化 ——')
set_font(run, size=16, color=(100, 100, 100))

for _ in range(3):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('【完整版】')
set_font(run, size=18, bold=True, color=(220, 38, 38))

for _ in range(6):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('面向：产品部 · 研发部 · 测试部 · 运营部')
set_font(run, size=12, color=(100, 100, 100))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('版本：V2.0（Hermes Agent 多轮对话版）')
set_font(run, size=12, color=(100, 100, 100))

doc.add_page_break()

# ========== 目录 ==========
add_heading(doc, '目录', level=1)
toc_items = [
    '一、项目背景与问题分析',
    '二、产品目标与核心价值',
    '三、产品架构设计（Hermes Agent 多轮对话架构）',
    '四、Hermes Agent 多轮对话优化方案设计',
    '五、完整提示词 Skills 模板库',
    '六、角色三视图设定专项设计',
    '七、主体设置页面交互设计（含用户确认）',
    '八、业务流程设计',
    '九、跨角色协作泳道设计',
    '十、技术实现方案',
    '十一、测试验收方案',
    '十二、项目排期与风险管理',
]
for item in toc_items:
    p = doc.add_paragraph()
    run = p.add_run(item)
    set_font(run, size=12)
    p.paragraph_format.space_after = Pt(6)

doc.add_page_break()

# ========== 第一章 ==========
add_heading(doc, '一、项目背景与问题分析', level=1)

add_heading(doc, '1.1 项目背景', level=2)
add_para(doc, '人工智能短剧创作系统是一款面向短剧创作者的全流程人工智能辅助工具，覆盖剧本生成、大纲拆解、主体资产管理、分镜生成、视频合成等核心环节。其中，主体资产管理模块（角色、场景、道具）是整个系统的视觉基础，其图片生成质量直接决定了最终短剧视频的视觉效果和用户体验。')
add_para(doc, '当前系统已实现主体图片的批量生成与单张重生成功能，但在实际使用过程中，用户反馈图片生成质量不稳定，尤其在人物角色图片方面存在多种典型问题，严重影响创作效率和成品质量。')
add_para(doc, '为解决上述问题，本方案提出基于 Hermes Agent 底层多轮对话的提示词智能优化方案。与传统基于规则的模板匹配不同，本方案借鉴短剧编剧和视频生成模块的成熟经验，通过 Hermes Agent 底层进行多轮隐式对话，反复迭代优化提示词，界面不显示对话内容，用户仅需确认最终优化结果。')

add_heading(doc, '1.2 现存问题详细分析', level=2)
add_para(doc, '通过对近一个月用户生成图片的抽样分析与用户反馈收集，我们梳理出以下四类高频问题：')

problems = [
    ['问题类型', '表现描述', '发生频率', '影响程度'],
    ['图片生成精度不高', '画面模糊、细节丢失、纹理不清晰、边缘锯齿明显，整体画质偏低', '约45%', '高'],
    ['人物图片色彩异常', '人物皮肤出现非自然色块、服饰颜色溢出、背景色污染主体、色彩断层', '约32%', '高'],
    ['斗鸡眼问题', '人物双眼视线不聚焦、瞳孔位置偏移、双眼朝向不一致，呈现斗鸡眼状态', '约28%', '极高'],
    ['人脸崩坏', '面部五官扭曲、脸型变形、五官比例失调、多脸或残脸现象', '约25%', '极高'],
]
add_table(doc, problems[0], problems[1:], col_widths=[3, 6, 2, 2])

add_heading(doc, '1.3 问题根因总结', level=2)
add_para(doc, '综合以上分析，当前图片生成质量问题的核心根因可归纳为以下三点：')
add_bullet(doc, '第一点：提示词结构过于简单。现有提示词仅包含基础描述加通用风格词，缺少结构化的质量控制层、特征锁定层、负面排除层，无法有效引导模型生成高质量图片。')
add_bullet(doc, '第二点：缺少迭代优化机制。当前提示词为一次性生成，缺少多轮迭代优化过程，无法根据模型反馈和问题特征进行针对性调整。')
add_bullet(doc, '第三点：无质量检测与自动优化闭环。生成后缺少自动质量检测环节，问题图片无法被自动识别和重试优化，全部依赖用户人工判断和手动重生成。')

doc.add_page_break()

# ========== 第二章 ==========
add_heading(doc, '二、产品目标与核心价值', level=1)

add_heading(doc, '2.1 产品目标', level=2)
goals = [
    ['目标维度', '具体指标', '当前基线', '目标值'],
    ['图片一次成功率', '首次生成即满足用户需求的比例', '约55%', '≥85%'],
    ['人脸崩坏率', '人物图片出现面部五官扭曲的比例', '约25%', '≤5%'],
    ['斗鸡眼发生率', '人物图片出现双眼视线异常的比例', '约28%', '≤3%'],
    ['色彩异常率', '图片出现非自然色块或颜色溢出的比例', '约32%', '≤8%'],
    ['用户重生成次数', '平均每个主体需要重生成的次数', '2.8次', '≤1.2次'],
    ['提示词优化耗时', '系统完成提示词优化的平均时间', '无优化', '≤8秒'],
]
add_table(doc, goals[0], goals[1:], col_widths=[3, 5, 2.5, 2.5])

add_heading(doc, '2.2 核心价值', level=2)
add_bullet(doc, '用户价值：对创作者，大幅提升图片生成质量与一次成功率，减少反复调整和重生成的时间成本，让创作者更专注于内容创作本身。')
add_bullet(doc, '产品价值：对产品，建立"Hermes Agent 多轮对话提示词优化引擎加质量检测闭环"的技术壁垒，提升产品核心竞争力和用户留存率。')
add_bullet(doc, '团队价值：对团队，通过 Hermes Agent 多轮对话的隐式优化机制和用户确认机制，降低生成结果的不可控性，减少售后问题和用户投诉。')
add_bullet(doc, '技术价值：对技术，探索 Hermes Agent 在视觉生成提示词优化领域的应用，为后续视频生成提示词优化、分镜脚本优化等场景积累经验。')

doc.add_page_break()

# ========== 第三章 ==========
add_heading(doc, '三、产品架构设计（Hermes Agent 多轮对话架构）', level=1)

add_heading(doc, '3.1 整体架构概述', level=2)
add_para(doc, '本方案在现有人工智能短剧系统架构基础上，引入 Hermes Agent 多轮对话引擎作为提示词优化的核心驱动，新增"隐式对话优化层"和"图像质量检测服务"两大核心模块，并在主体设置页面增加"用户确认"交互环节，形成"用户输入到 Hermes Agent 多轮隐式优化到用户确认到质量检测到反馈迭代"的完整闭环。')
add_para(doc, '与传统基于规则的模板匹配不同，Hermes Agent 多轮对话优化的核心特点是：界面不显示对话内容，所有优化过程在 Hermes 底层通过多轮隐式对话完成，用户仅看到最终优化结果并进行确认。')

add_heading(doc, '3.2 产品架构图', level=2)
add_para(doc, '图3-1  基于 Hermes Agent 多轮对话的提示词优化产品架构图', indent=False)

# 架构图（用表格模拟）
add_flow_box(doc, '用户层', '创作者 / 运营人员 / 质检人员', (30, 58, 138))
add_arrow(doc)
add_flow_box(doc, '应用层', '主体设置模块（含用户确认） | 图片生成模块 | 质量检测与反馈模块', (37, 99, 235))
add_arrow(doc)
add_flow_box(doc, 'Hermes Agent 多轮对话层（核心新增）', 'Skills 调度器 | 多轮对话管理器 | 提示词优化 Agent | 质量分析 Agent | 迭代优化 Agent', (124, 58, 237))
add_arrow(doc)
add_flow_box(doc, '服务层', '提示词 Skills 模板库 | 人物特征锁定服务 | 模型网关 | 图像质量分析服务 | 用户确认与反馈服务', (16, 185, 129))
add_arrow(doc)
add_flow_box(doc, '数据层', '主体资产库 | 提示词知识库 | 生成历史库 | 用户行为日志 | 对话上下文存储', (217, 119, 6))

add_heading(doc, '3.3 架构分层说明', level=2)

add_heading(doc, '3.3.1 用户层', level=3)
add_bullet(doc, '创作者/运营人员：使用主体设置模块进行角色/场景/道具的创建、编辑和图片生成，是提示词优化功能的主要使用者。')
add_bullet(doc, '审核/质检人员：对生成图片进行质量抽检和审核，反馈问题案例用于 Hermes Agent Skills 模板的持续优化。')
add_bullet(doc, '系统管理员：维护提示词 Skills 模板库、质量检测规则和 Hermes Agent 配置。')

add_heading(doc, '3.3.2 应用层', level=3)
add_bullet(doc, '主体设置模块：升级现有主体设置页面，增加提示词预览、编辑和用户确认交互。')
add_bullet(doc, 'Hermes Agent 调用模块：负责将用户输入传递给 Hermes Agent，并接收优化后的提示词结果。')
add_bullet(doc, '图片生成模块：在现有生成能力基础上，接入 Hermes Agent 优化后的提示词和多模型调度策略。')
add_bullet(doc, '质量检测与反馈：新增模块，对生成图片进行自动质量检测，并支持用户人工确认反馈。')

add_heading(doc, '3.3.3 Hermes Agent 多轮对话层（核心新增）', level=3)
add_bullet(doc, 'Skills 调度器：根据主体类型（角色/场景/道具）和用户需求，调度对应的提示词优化 Skills 模板。')
add_bullet(doc, '多轮对话管理器：管理 Hermes Agent 底层的多轮隐式对话上下文，控制对话轮次和终止条件。')
add_bullet(doc, '提示词优化 Agent：Hermes Agent 的核心角色，负责通过多轮对话逐步完善提示词结构。')
add_bullet(doc, '质量分析 Agent：分析潜在的质量问题，生成针对性的负面提示词和预防措施。')
add_bullet(doc, '迭代优化 Agent：根据质量检测反馈，迭代优化提示词，形成闭环。')

add_heading(doc, '3.3.4 服务层', level=3)
add_bullet(doc, '提示词 Skills 模板库：存储各类主体的提示词优化 Skills 模板，包括角色、场景、道具等。')
add_bullet(doc, '人物特征锁定服务：从用户描述中提取并标准化人物特征（面部、肤色、发型、服饰等）。')
add_bullet(doc, '模型网关：统一调度多个图片生成模型，支持按场景/主体类型智能路由。')
add_bullet(doc, '图像质量分析服务：集成人脸检测、色彩分析、清晰度评估、眼部检测等算法。')
add_bullet(doc, '用户确认与反馈服务：管理用户确认状态、反馈数据收集和提示词优化建议。')

add_heading(doc, '3.3.5 数据层', level=3)
add_bullet(doc, '主体资产库：存储角色、场景、道具的基础信息、参考图和生成图片历史。')
add_bullet(doc, '提示词知识库：存储 Hermes Agent Skills 模板、优化规则、最佳实践案例。')
add_bullet(doc, '生成历史库：记录每次生成的图片、提示词参数、模型版本、质量评分。')
add_bullet(doc, '用户行为日志：记录用户确认、修改、重生成、采纳等行为数据。')
add_bullet(doc, '对话上下文存储：存储 Hermes Agent 多轮对话的上下文信息，支持断点续优化。')

doc.add_page_break()

# ========== 第四章 ==========
add_heading(doc, '四、Hermes Agent 多轮对话优化方案设计', level=1)

add_heading(doc, '4.1 设计理念', level=2)
add_para(doc, '本方案借鉴短剧编剧模块（骨架生成器、大纲填充器、主模板）和视频生成模块（分镜拆解、镜头拆解）的成熟多轮对话经验，将 Hermes Agent 多轮对话机制引入图片生成提示词优化领域。')
add_para(doc, '核心设计原则：界面不显示对话内容，所有优化过程在 Hermes 底层通过多轮隐式对话完成。用户仅需输入基础描述，系统通过 Hermes Agent 进行多轮自我提问与回答，逐步完善提示词，最终输出优化结果供用户确认。')

add_heading(doc, '4.2 Hermes Agent 多轮对话流程', level=2)
add_para(doc, 'Hermes Agent 接收用户输入的主体信息后，按以下八步流程完成多轮隐式对话优化：')

flow_steps = [
    ['轮次', '对话角色', '对话内容（隐式，界面不显示）', '输出结果'],
    ['第1轮', '提示词优化 Agent', '解析用户输入的主体名称、描述、参考图、风格偏好等信息，识别主体类型', '结构化的主体信息摘要'],
    ['第2轮', '特征提取 Agent', '从描述文本中提取人物特征（年龄/性别/发型/服饰/肤色等）或场景/道具特征', '标准化特征清单'],
    ['第3轮', 'Skills 匹配 Agent', '根据主体类型加风格加镜头等标签，从 Skills 模板库中匹配最佳优化模板', '匹配的 Skills 模板 ID'],
    ['第4轮', '动态填充 Agent', '将提取的特征填入 Skills 模板占位符，生成初始正面提示词', '初始正面提示词 V1'],
    ['第5轮', '质量增强 Agent', '根据主体类型自动注入质量增强词和问题预防词，针对斗鸡眼、人脸崩坏等问题', '增强后正面提示词 V2'],
    ['第6轮', '负面词组合 Agent', '按优先级组合负面提示词，控制总长度在合理范围', '负面提示词 V1'],
    ['第7轮', '校验优化 Agent', '检查关键词冲突、重复、长度超限，进行去重和优化排序', '最终正面+负面提示词'],
    ['第8轮', '结果输出 Agent', '输出正面提示词、负面提示词、推荐生成参数，供前端预览和用户确认', '完整优化结果包'],
]
add_table(doc, flow_steps[0], flow_steps[1:], col_widths=[1.5, 2.5, 6, 3])

add_heading(doc, '4.3 多轮对话状态机', level=2)
add_para(doc, 'Hermes Agent 多轮对话采用状态机管理，确保每轮对话有明确的输入、输出和终止条件：')

states = [
    ['状态', '触发条件', '执行动作', '下一状态'],
    ['初始化', '用户点击优化按钮', '接收用户输入，创建对话上下文', '信息解析'],
    ['信息解析', '上下文创建完成', '第1轮对话：解析主体信息', '特征提取'],
    ['特征提取', '信息解析完成', '第2轮对话：提取标准化特征', 'Skills匹配'],
    ['Skills匹配', '特征提取完成', '第3轮对话：匹配最佳模板', '动态填充'],
    ['动态填充', 'Skills匹配完成', '第4轮对话：填充模板生成初始提示词', '质量增强'],
    ['质量增强', '初始提示词生成', '第5轮对话：注入质量增强词', '负面词组合'],
    ['负面词组合', '质量增强完成', '第6轮对话：组合负面提示词', '校验优化'],
    ['校验优化', '负面词组合完成', '第7轮对话：校验和优化排序', '结果输出'],
    ['结果输出', '校验优化完成', '第8轮对话：输出最终结果', '等待用户确认'],
    ['用户确认', '用户点击确认', '保存优化结果，触发生成', '结束'],
    ['用户修改', '用户修改提示词', '记录修改，更新对话上下文', '结果输出'],
    ['迭代优化', '质量检测不通过', '基于反馈重新进入质量增强', '质量增强'],
]
add_table(doc, states[0], states[1:], col_widths=[2, 3, 5, 2.5])

add_heading(doc, '4.4 隐式对话的用户体验设计', level=2)
add_para(doc, '由于界面不显示对话内容，为确保用户体验，采用以下设计：')
add_bullet(doc, '加载动画：优化过程中显示精美的加载动画，提示"正在通过 Hermes Agent 智能优化提示词..."。')
add_bullet(doc, '进度提示：分阶段显示优化进度（信息解析、特征提取、模板匹配、质量增强...），让用户感知优化过程。')
add_bullet(doc, '耗时控制：单主体优化控制在 8 秒以内，超过 10 秒自动降级为快速优化模式。')
add_bullet(doc, '结果展示：优化完成后，以对比卡片形式展示原始描述与优化后提示词，支持一键采纳或手动编辑。')

doc.add_page_break()

# ========== 第五章 ==========
add_heading(doc, '五、完整提示词 Skills 模板库', level=1)

add_heading(doc, '5.1 Skills 模板库概述', level=2)
add_para(doc, '提示词 Skills 模板库是 Hermes Agent 多轮对话优化的知识基础，参考短剧编剧模块的三模板架构（骨架生成器、大纲填充器、主模板）和视频生成模块的分镜拆解模板，设计为分层、可扩展的 Skills 体系。')
add_para(doc, '每个 Skills 模板包含：系统角色定义、多轮对话流程、输出格式规范、质量控制规则四个部分。Hermes Agent 根据主体类型自动调度对应的 Skills 模板。')

add_heading(doc, '5.2 角色提示词 Skills 模板', level=2)

add_heading(doc, '5.2.1 角色基础设定 Skills', level=3)
add_para(doc, '【系统角色】你是一位专业的人物形象设计提示词优化专家，精通 Stable Diffusion、Midjourney 等图像生成模型的提示词工程，擅长通过多轮自我对话完善人物描述。', indent=False)
add_para(doc, '【多轮对话流程】', indent=False)
add_bullet(doc, '第1轮：解析用户输入的人物名称和基础描述，识别性别、年龄区间、身份职业。')
add_bullet(doc, '第2轮：提取面部特征（脸型、五官、眼睛、鼻子、嘴巴、肤色），标准化为描述性词语。')
add_bullet(doc, '第3轮：提取发型和发色特征，补充发型细节（长度、卷曲度、刘海样式）。')
add_bullet(doc, '第4轮：提取服饰特征，补充服装款式、颜色、材质、配饰。')
add_bullet(doc, '第5轮：根据人物身份和场景，补充姿态、表情、眼神方向。')
add_bullet(doc, '第6轮：注入质量增强词和人物专用负面提示词。')

add_heading(doc, '5.2.2 角色三视图设定 Skills', level=3)
add_para(doc, '【系统角色】你是一位专业的角色三视图设计专家，精通角色正面、侧面、背面三视图的标准化描述。', indent=False)
add_para(doc, '【输出格式】三视图提示词必须包含：正面视图（front view）、侧面视图（side view）、背面视图（back view），统一服装、发型、配饰，确保三视图一致性。', indent=False)
add_para(doc, '【质量控制】强制包含：same character, consistent outfit, consistent hairstyle, character sheet, multiple views, turnaround。', indent=False)

add_heading(doc, '5.2.3 角色一致性锁定 Skills', level=3)
add_para(doc, '【系统角色】你是一位人物一致性专家，擅长在多镜头、多场景中保持人物形象的高度一致。', indent=False)
add_para(doc, '【多轮对话流程】', indent=False)
add_bullet(doc, '第1轮：分析是否为多镜头场景，识别镜头切换关键词（先是、然后、接着、镜头切换）。')
add_bullet(doc, '第2轮：提取人物核心特征指纹（面部特征+服饰特征+标志性配饰）。')
add_bullet(doc, '第3轮：为每个镜头生成独立提示词，共享核心特征指纹。')
add_bullet(doc, '第4轮：注入一致性约束词（same character, consistent appearance, identical face）。')

add_heading(doc, '5.3 场景提示词 Skills 模板', level=2)

add_heading(doc, '5.3.1 室内场景 Skills', level=3)
add_para(doc, '【系统角色】你是一位专业的室内场景设计提示词优化专家，精通各类室内空间的视觉化描述。', indent=False)
add_para(doc, '【覆盖场景】办公室、会议室、咖啡厅、餐厅、酒吧、家居客厅、卧室、厨房、医院、学校、商场、电梯、走廊、卫生间等22类室内场景。', indent=False)
add_para(doc, '【光影标准】每类场景内置标准光影模板，例如：办公室（明亮自然光，冷白色调，大面积窗户）、咖啡厅（暖黄色调，柔和壁灯，木质纹理）、酒吧（低照度，霓虹氛围灯，深色背景）。', indent=False)

add_heading(doc, '5.3.2 室外场景 Skills', level=3)
add_para(doc, '【系统角色】你是一位专业的室外场景设计提示词优化专家，精通自然景观和城市景观的视觉化描述。', indent=False)
add_para(doc, '【覆盖场景】城市街道、公园、广场、海边、山林、田野、沙漠、雪地、雨天、夜景、黄昏、清晨等18类室外场景。', indent=False)
add_para(doc, '【光影标准】每类场景内置时间与天气对应的光影模板，确保场景氛围准确。', indent=False)

add_heading(doc, '5.4 道具提示词 Skills 模板', level=2)
add_para(doc, '【系统角色】你是一位专业的道具设计提示词优化专家，精通各类物品的精细化描述。', indent=False)
add_para(doc, '【覆盖道具】武器（刀剑、枪械）、电子产品（手机、电脑）、交通工具（汽车、自行车）、日常用品（杯具、文具）、首饰、乐器、书籍、食品等15类道具。', indent=False)
add_para(doc, '【质量控制】道具提示词强制包含：材质描述、光影反射、背景处理（纯色背景或场景融合）、细节纹理。', indent=False)

add_heading(doc, '5.5 通用质量增强 Skills 模板', level=2)
add_para(doc, '【系统角色】你是一位图像质量控制专家，精通各类质量问题的预防和优化。', indent=False)
add_para(doc, '【正向质量词库】masterpiece, best quality, ultra detailed, 8k resolution, photorealistic, cinematic lighting, sharp focus, depth of field, professional photography, high dynamic range。', indent=False)
add_para(doc, '【负向质量词库】low quality, blurry, distorted, deformed, disfigured, bad anatomy, bad proportions, extra limbs, missing limbs, fused fingers, too many fingers, long neck, cross-eyed, asymmetrical eyes, color fringing, color bleeding, watermark, text, signature。', indent=False)
add_para(doc, '【问题专项预防】针对斗鸡眼：symmetrical eyes, focused gaze, normal eye position, detailed pupils；针对人脸崩坏：detailed face, perfect face, symmetrical facial features, realistic skin texture。', indent=False)

add_heading(doc, '5.6 Skills 模板库统计', level=2)
skills_stats = [
    ['Skills 分类', '模板数量', '适用主体', '优化轮次'],
    ['角色基础设定', '12套', '人物角色', '6轮'],
    ['角色三视图设定', '3套', '人物角色', '5轮'],
    ['角色一致性锁定', '4套', '人物角色', '4轮'],
    ['室内场景', '22套', '场景', '5轮'],
    ['室外场景', '18套', '场景', '5轮'],
    ['道具设计', '15套', '道具', '4轮'],
    ['通用质量增强', '8套', '全部', '3轮'],
    ['风格迁移', '10套', '全部', '4轮'],
    ['合计', '92套', '-', '-'],
]
add_table(doc, skills_stats[0], skills_stats[1:], col_widths=[3, 2, 3, 2])

doc.add_page_break()

# ========== 第六章 ==========
add_heading(doc, '六、角色三视图设定专项设计', level=1)

add_heading(doc, '6.1 三视图设定的重要性', level=2)
add_para(doc, '角色三视图（正面、侧面、背面）是短剧创作中保持人物一致性的关键资产。通过三视图设定，系统可以在后续分镜生成中准确还原人物形象，避免不同镜头中人物形象不一致的问题。')

add_heading(doc, '6.2 三视图录入流程', level=2)
add_para(doc, '用户可通过以下两种方式录入角色三视图：')

add_heading(doc, '6.2.1 方式一：Hermes Agent 智能生成三视图', level=3)
add_bullet(doc, '步骤1：用户在主体设置页面选择"角色"类型，填写角色基础描述。')
add_bullet(doc, '步骤2：点击"生成三视图"按钮，触发 Hermes Agent 三视图设定 Skills。')
add_bullet(doc, '步骤3：Hermes Agent 通过5轮隐式对话，生成标准化的三视图提示词。')
add_bullet(doc, '步骤4：系统调用图片生成模型，生成包含正面、侧面、背面的三视图图片。')
add_bullet(doc, '步骤5：用户确认三视图图片，系统自动保存为角色参考图。')

add_heading(doc, '6.2.2 方式二：手动上传三视图参考图', level=3)
add_bullet(doc, '步骤1：用户在主体设置页面选择"角色"类型。')
add_bullet(doc, '步骤2：点击"上传三视图"按钮，分别上传正面、侧面、背面三张图片。')
add_bullet(doc, '步骤3：Hermes Agent 自动分析三张图片，提取人物特征指纹。')
add_bullet(doc, '步骤4：系统生成标准化的人物描述，供用户确认和编辑。')

add_heading(doc, '6.3 三视图 Hermes Agent 对话流程', level=2)
three_view_flow = [
    ['轮次', '对话内容', '输出'],
    ['第1轮', '解析角色基础描述，识别性别、年龄、身份', '角色基础信息'],
    ['第2轮', '提取面部特征和发型特征', '头部特征清单'],
    ['第3轮', '提取服饰和配饰特征，确保三视图一致', '服饰特征清单'],
    ['第4轮', '生成三视图专用提示词，包含正面、侧面、背面视图', '三视图提示词'],
    ['第5轮', '注入一致性约束和质量增强词', '最终三视图提示词'],
]
add_table(doc, three_view_flow[0], three_view_flow[1:], col_widths=[1.5, 7, 3])

add_heading(doc, '6.4 三视图提示词模板示例', level=2)
add_para(doc, '【正面视图】character sheet, front view, 1girl, 25 years old, long black hair, blue eyes, white blouse, black skirt, standing, neutral expression, looking at viewer, full body, white background, masterpiece, best quality, ultra detailed', indent=False)
add_para(doc, '【侧面视图】character sheet, side view, 1girl, 25 years old, long black hair, blue eyes, white blouse, black skirt, standing, profile view, full body, white background, masterpiece, best quality, ultra detailed', indent=False)
add_para(doc, '【背面视图】character sheet, back view, 1girl, 25 years old, long black hair, white blouse, black skirt, standing, back of head, full body, white background, masterpiece, best quality, ultra detailed', indent=False)
add_para(doc, '【一致性约束】same character, consistent outfit, consistent hairstyle, identical accessories, character turnaround, multiple views in one image', indent=False)

doc.add_page_break()

# ========== 第七章 ==========
add_heading(doc, '七、主体设置页面交互设计（含用户确认）', level=1)

add_heading(doc, '7.1 页面整体布局', level=2)
add_para(doc, '主体设置页面采用左右分栏布局：左侧为主体列表（角色、场景、道具三个标签页），右侧为选中主体的编辑区域，包含基础信息、参考图、提示词优化结果、生成参数等模块。')

add_heading(doc, '7.2 顶部工具栏设计', level=2)
add_para(doc, '页面顶部固定工具栏，从左到右依次为：')
add_bullet(doc, '页面标题：主体设置（第X集）')
add_bullet(doc, '资产库按钮：点击打开资产库选择弹窗，可从已有资产中导入主体。')
add_bullet(doc, '添加主体下拉按钮：点击展开下拉菜单，可选择添加角色、添加场景、添加道具。')
add_bullet(doc, '分辨率选择：支持1K、1.5K、2K、3K、4K分辨率选择。')
add_bullet(doc, '批量下载按钮：批量下载所有主体图片。')
add_bullet(doc, '图片模型选择器：选择用于生成图片的AI模型。')

add_heading(doc, '7.3 提示词优化结果展示', level=2)
add_para(doc, '当用户输入主体描述并点击"Hermes 智能优化"按钮后，系统展示优化结果：')
add_bullet(doc, '优化概览卡片：显示命中的 Skills 模板、优化轮次、预计质量提升。')
add_bullet(doc, '原始描述与优化后提示词对比：左右分栏展示，高亮新增和修改的关键词。')
add_bullet(doc, '负面提示词展示：以标签形式展示负面提示词，支持手动删除和添加。')
add_bullet(doc, '推荐生成参数：步数、相关性系数、采样器等参数的推荐值。')

add_heading(doc, '7.4 用户确认操作详细设计', level=2)
add_para(doc, '用户确认操作是本次方案的核心交互点，确保用户在生成前明确知晓并认可 Hermes Agent 优化的提示词。')

add_heading(doc, '7.4.1 确认按钮状态与逻辑', level=3)
add_bullet(doc, '默认状态：按钮文案为「确认并生成图片」，主按钮样式（靛蓝渐变），可点击。')
add_bullet(doc, '未填写必填信息时：按钮禁用，悬浮提示"请完善主体描述信息"。')
add_bullet(doc, 'Hermes 优化中：按钮变为加载状态，文案"Hermes 优化中..."，禁止重复点击。')
add_bullet(doc, '用户修改过提示词时：按钮旁显示「已自定义」标签，提示用户当前使用的是自定义提示词。')
add_bullet(doc, '生成中：按钮变为加载状态，文案"生成中..."，禁止重复点击。')

add_heading(doc, '7.4.2 确认前校验', level=3)
add_bullet(doc, '主体描述非空校验：确保有足够的描述信息用于生成。')
add_bullet(doc, '提示词长度校验：正面提示词建议五十至三百词，负面提示词建议二十至一百词，超限时给出警告但不强制阻断。')
add_bullet(doc, '敏感词校验：检查提示词中是否包含违规内容。')
add_bullet(doc, '特征完整性提醒：人物角色缺少关键特征（如性别/年龄）时，弹出确认框"检测到特征信息不完整，是否继续生成？"。')

add_heading(doc, '7.4.3 二次确认场景', level=3)
add_bullet(doc, '用户从未修改过提示词，首次点击生成时：弹出轻量确认"Hermes Agent 已为您智能优化提示词，确认使用该提示词生成图片？"，含「确认生成」和「再看看」按钮。')
add_bullet(doc, '用户选择高分辨率或多图生成时：提示"高分辨率/多图生成将消耗更多时间和资源，确认继续？"。')

doc.add_page_break()

# ========== 第八章 ==========
add_heading(doc, '八、业务流程设计', level=1)

add_heading(doc, '8.1 主体创建与图片生成主流程', level=2)
add_para(doc, '图8-1  主体创建与图片生成主流程图', indent=False)

main_flow = [
    '用户进入主体设置页面',
    '选择添加主体类型（角色/场景/道具）或从资产库导入',
    '填写主体基础信息（名称、描述、参考图）',
    '点击「Hermes 智能优化」按钮',
    'Hermes Agent 多轮隐式对话优化（8轮）',
    '展示优化结果（正面提示词+负面提示词+推荐参数）',
    '用户确认或手动编辑提示词',
    '点击「确认并生成图片」',
    '调用图片生成模型生成图片',
    '自动质量检测（人脸/色彩/清晰度/眼部）',
    '质量达标 → 保存图片，完成',
    '质量不达标 → Hermes Agent 迭代优化，重新生成（最多3次）',
]
for step in main_flow:
    add_flow_box(doc, step, '', (37, 99, 235))
    if step != main_flow[-1]:
        add_arrow(doc)

add_heading(doc, '8.2 快捷创作流程（无剧本模式）', level=2)
add_para(doc, '支持用户无需上传剧本，直接创建主体并生成分镜视频：')
add_bullet(doc, '步骤1：用户在首页选择「快捷创作」模式，直接进入主体设置页面。')
add_bullet(doc, '步骤2：用户手动添加或从资产库选择角色、场景、道具。')
add_bullet(doc, '步骤3：每个主体通过 Hermes Agent 优化提示词并生成图片。')
add_bullet(doc, '步骤4：用户进入分镜脚本页面，手动创建分镜。')
add_bullet(doc, '步骤5：分镜脚本通过 Hermes Agent 优化后生成视频。')

doc.add_page_break()

# ========== 第九章 ==========
add_heading(doc, '九、跨角色协作泳道设计', level=1)

add_heading(doc, '9.1 泳道图概述', level=2)
add_para(doc, '图9-1  提示词优化跨角色协作泳道图', indent=False)
add_para(doc, '本流程涉及五个核心角色：用户、前端应用、Hermes Agent、图片生成服务、质量检测服务。', indent=False)

swimlanes = [
    ['用户', '前端应用', 'Hermes Agent', '图片生成服务', '质量检测服务'],
    ['填写主体描述', '', '', '', ''],
    ['点击优化按钮', '接收请求，显示加载动画', '', '', ''],
    ['', '调用 Hermes Agent 接口', '开始多轮对话优化', '', ''],
    ['', '', '第1-8轮隐式对话', '', ''],
    ['', '接收优化结果', '返回优化后提示词', '', ''],
    ['', '展示优化结果对比', '', '', ''],
    ['确认并生成', '提交生成请求', '', '', ''],
    ['', '', '', '生成图片', ''],
    ['', '接收图片', '', '返回图片', ''],
    ['', '提交质量检测', '', '', '执行检测'],
    ['', '接收检测结果', '', '', '返回报告'],
    ['查看生成结果', '达标则保存，不达标则触发迭代', '', '', ''],
]
add_table(doc, swimlanes[0], swimlanes[1:], col_widths=[3, 3, 3, 3, 3])

doc.add_page_break()

# ========== 第十章 ==========
add_heading(doc, '十、技术实现方案', level=1)

add_heading(doc, '10.1 技术栈', level=2)
tech_stack = [
    ['层级', '技术选型', '说明'],
    ['前端', 'Vue3 + TypeScript + Element Plus', '主体设置页面、用户确认交互'],
    ['Agent层', 'Hermes Agent + Skills 机制', '多轮隐式对话优化核心'],
    ['后端', 'Node.js + Python 微服务', 'Hermes Agent 调度、接口服务'],
    ['图像生成', 'Stable Diffusion / 自研模型', '图片生成引擎'],
    ['质量检测', 'OpenCV + 人脸检测模型', '人脸、色彩、清晰度、眼部检测'],
    ['数据存储', 'MySQL + Redis + 对象存储', '业务数据、缓存、图片存储'],
    ['消息队列', 'RabbitMQ', '异步生成任务调度'],
]
add_table(doc, tech_stack[0], tech_stack[1:], col_widths=[2, 5, 5])

add_heading(doc, '10.2 Hermes Agent 集成方案', level=2)
add_bullet(doc, 'Skills 注册：将提示词优化 Skills 模板注册到 Hermes Agent，按主体类型分类。')
add_bullet(doc, '对话上下文管理：使用 Redis 存储多轮对话上下文，支持超时自动清理。')
add_bullet(doc, '流式输出：Hermes Agent 优化过程采用流式输出，前端实时展示进度。')
add_bullet(doc, '降级策略：Hermes Agent 不可用时，自动降级为基于规则的快速优化模式。')
add_bullet(doc, '监控告警：监控 Hermes Agent 响应时间、成功率、优化轮次等指标。')

add_heading(doc, '10.3 接口设计', level=2)
add_para(doc, '核心接口：POST /api/v1/prompt/optimize', indent=False)
add_bullet(doc, '请求参数：subject_type（角色/场景/道具）、name、description、reference_images、style_preference。')
add_bullet(doc, '响应参数：enhanced_prompt（正面提示词）、negative_prompt（负面提示词）、recommended_params（推荐参数）、optimization_trace（优化轨迹，用于调试）。')
add_bullet(doc, '超时设置：默认15秒，支持客户端自定义。')

doc.add_page_break()

# ========== 第十一章 ==========
add_heading(doc, '十一、测试验收方案', level=1)

add_heading(doc, '11.1 测试策略', level=2)
add_bullet(doc, '单元测试：覆盖 Hermes Agent 各轮对话逻辑、Skills 模板匹配、提示词校验。')
add_bullet(doc, '集成测试：覆盖前端到 Hermes Agent 到图片生成的完整链路。')
add_bullet(doc, '性能测试：验证单主体优化耗时≤8秒，并发100用户时响应时间≤15秒。')
add_bullet(doc, '质量测试：使用1000张测试图片，验证人脸崩坏率≤5%、斗鸡眼率≤3%。')
add_bullet(doc, '用户验收测试：邀请20位真实创作者试用，收集满意度反馈。')

add_heading(doc, '11.2 验收标准', level=2)
acceptance = [
    ['验收项', '标准', '测试方法'],
    ['Hermes 优化功能', '点击优化按钮后8秒内返回结果', '功能测试+性能测试'],
    ['提示词质量', '优化后提示词包含质量增强词和负面词', '人工审核+自动化检查'],
    ['用户确认交互', '确认按钮状态、校验、二次确认均正常', '功能测试'],
    ['图片质量', '人脸崩坏率≤5%，斗鸡眼率≤3%', '批量生成+人工审核'],
    ['三视图功能', '三视图生成一致性≥90%', '人工审核'],
    ['降级策略', 'Hermes 不可用时自动降级', '故障注入测试'],
]
add_table(doc, acceptance[0], acceptance[1:], col_widths=[3, 5, 4])

doc.add_page_break()

# ========== 第十二章 ==========
add_heading(doc, '十二、项目排期与风险管理', level=1)

add_heading(doc, '12.1 项目排期', level=2)
schedule = [
    ['阶段', '任务', '工期', '负责人'],
    ['第一阶段', '需求评审与技术方案确认', '1周', '产品+技术'],
    ['第二阶段', 'Hermes Agent Skills 模板开发', '2周', '算法+后端'],
    ['第三阶段', '前端主体设置页面升级', '2周', '前端'],
    ['第四阶段', '质量检测服务开发', '1.5周', '算法'],
    ['第五阶段', '集成测试与联调', '1周', '测试+全栈'],
    ['第六阶段', '灰度发布与用户反馈收集', '1周', '运营+产品'],
    ['第七阶段', '全量发布与持续优化', '持续', '全体'],
    ['合计', '', '8.5周', ''],
]
add_table(doc, schedule[0], schedule[1:], col_widths=[2, 5, 2, 3])

add_heading(doc, '12.2 风险管理', level=2)
risks = [
    ['风险项', '风险等级', '应对措施'],
    ['Hermes Agent 响应时间过长', '高', '设置超时降级，优化 Skills 模板减少轮次'],
    ['Hermes Agent 优化效果不稳定', '高', '建立 A/B 测试机制，持续优化 Skills 模板'],
    ['质量检测准确率不足', '中', '积累标注数据，持续训练检测模型'],
    ['用户不接受新交互流程', '中', '提供开关可切换新旧模式，收集反馈逐步优化'],
    ['并发量过大导致服务不稳定', '中', '限流+排队机制，弹性扩容'],
]
add_table(doc, risks[0], risks[1:], col_widths=[4, 2, 6])

# 保存文档
output_path = r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\人工智能短剧系统_图片生成提示词优化产品方案_HermesAgent完整版.docx'
doc.save(output_path)
print(f'文档已生成：{output_path}')
