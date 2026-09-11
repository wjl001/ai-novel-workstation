#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成基于 Hermes Agent 多轮对话的图片生成提示词优化产品方案 - 领导汇报版"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_font(run, name='微软雅黑', size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    sizes = {1: 18, 2: 15, 3: 13, 4: 12}
    set_font(run, size=sizes.get(level, 12), bold=True, color=(30, 58, 138))
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_para(doc, text, size=11, bold=False, indent=True, color=None):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(22)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    return p

def add_bullet(doc, text, size=11, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(22 + level * 22)
    run = p.add_run('● ' + text)
    set_font(run, size=size)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(2)
    return p

def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_font(run, size=10, bold=True, color=(255, 255, 255))
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), '1E3A8A')
        cell._tc.get_or_add_tcPr().append(shading)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_font(run, size=10)
            if r_idx % 2 == 1:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'F0F4FF')
                cell._tc.get_or_add_tcPr().append(shading)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    return table

def add_flow_box(doc, title, content='', color=(30, 58, 138)):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_font(run, size=11, bold=True, color=(255, 255, 255))
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), '%02X%02X%02X' % color)
    cell._tc.get_or_add_tcPr().append(shading)
    if content:
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = p2.add_run(content)
        set_font(run2, size=9, color=(255, 255, 255))
    return table

def add_arrow(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('↓')
    set_font(run, size=16, bold=True, color=(30, 58, 138))

# 创建文档
doc = Document()
style = doc.styles['Normal']
style.font.name = '微软雅黑'
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

# ========== 封面 ==========
for _ in range(5):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('人工智能短剧创作系统')
set_font(run, size=26, bold=True, color=(30, 58, 138))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('图片生成提示词优化产品方案')
set_font(run, size=22, bold=True, color=(30, 58, 138))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('—— 基于 Hermes Agent 多轮对话智能优化 ——')
set_font(run, size=14, color=(100, 100, 100))

for _ in range(3):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('【领导汇报版】')
set_font(run, size=18, bold=True, color=(220, 38, 38))

for _ in range(8):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('版本：V2.0（Hermes Agent 多轮对话版）')
set_font(run, size=12, color=(100, 100, 100))

doc.add_page_break()

# ========== 一、项目背景与目标 ==========
add_heading(doc, '一、项目背景与目标', level=1)

add_heading(doc, '1.1 现存问题', level=2)
add_para(doc, '当前系统主体图片生成质量不稳定，用户反馈四类高频问题：图片精度不高（45%）、人物色彩异常（32%）、斗鸡眼（28%）、人脸崩坏（25%），导致平均每个主体需重生成2.8次，严重影响创作效率。')

add_heading(doc, '1.2 核心目标', level=2)
goals = [
    ['指标', '当前', '目标'],
    ['图片一次成功率', '55%', '≥85%'],
    ['人脸崩坏率', '25%', '≤5%'],
    ['斗鸡眼发生率', '28%', '≤3%'],
    ['平均重生成次数', '2.8次', '≤1.2次'],
]
add_table(doc, goals[0], goals[1:], col_widths=[5, 3, 3])

add_heading(doc, '1.3 方案核心', level=2)
add_para(doc, '引入 Hermes Agent 多轮对话引擎，借鉴短剧编剧和视频生成模块的成熟经验，通过底层多轮隐式对话迭代优化提示词。界面不显示对话内容，用户仅需确认最终优化结果，形成"输入到 Hermes 优化到用户确认到质量检测到反馈迭代"的完整闭环。')

doc.add_page_break()

# ========== 二、产品架构设计（重点） ==========
add_heading(doc, '二、产品架构设计（重点）', level=1)

add_heading(doc, '2.1 整体架构图', level=2)
add_para(doc, '图2-1  基于 Hermes Agent 多轮对话的提示词优化产品架构图', indent=False)
# 插入真正的产品架构图
doc.add_picture(r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\architecture_diagram.png', width=Inches(6.0))
last_paragraph = doc.paragraphs[-1]
last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

add_heading(doc, '2.2 核心新增层：Hermes Agent 多轮对话层', level=2)
add_para(doc, '这是本方案的核心创新点，包含五个关键组件：')
add_bullet(doc, 'Skills 调度器：根据主体类型（角色/场景/道具）自动调度对应的优化 Skills 模板，目前共92套模板。')
add_bullet(doc, '多轮对话管理器：管理底层隐式对话上下文，控制对话轮次（标准8轮）和终止条件，单主体优化≤8秒。')
add_bullet(doc, '提示词优化 Agent：核心角色，通过多轮自我提问与回答，逐步完善提示词结构。')
add_bullet(doc, '质量分析 Agent：针对斗鸡眼、人脸崩坏、色彩异常等问题，生成针对性的负面提示词。')
add_bullet(doc, '迭代优化 Agent：根据质量检测反馈，自动迭代优化提示词，最多3轮重试。')

add_heading(doc, '2.3 架构优势', level=2)
add_bullet(doc, '隐式优化：界面不显示对话内容，用户体验简洁，仅看到最终优化结果。')
add_bullet(doc, '知识复用：Skills 模板库可持续积累最佳实践，优化效果随使用不断提升。')
add_bullet(doc, '闭环迭代：质量检测结果自动反馈给 Hermes Agent，形成自优化闭环。')
add_bullet(doc, '降级保障：Hermes Agent 不可用时自动降级为规则优化，确保服务可用性。')

doc.add_page_break()

# ========== 三、Hermes Agent 多轮对话流程 ==========
add_heading(doc, '三、Hermes Agent 多轮对话流程', level=1)

add_heading(doc, '3.1 八轮隐式对话优化流程', level=2)
flow_steps = [
    ['轮次', 'Agent 角色', '核心动作', '输出'],
    ['第1轮', '提示词优化 Agent', '解析主体信息，识别类型', '主体信息摘要'],
    ['第2轮', '特征提取 Agent', '提取人物/场景/道具特征', '标准化特征清单'],
    ['第3轮', 'Skills 匹配 Agent', '匹配最佳优化模板', 'Skills 模板 ID'],
    ['第4轮', '动态填充 Agent', '填充模板生成初始提示词', '正面提示词 V1'],
    ['第5轮', '质量增强 Agent', '注入质量增强词和问题预防词', '正面提示词 V2'],
    ['第6轮', '负面词组合 Agent', '按优先级组合负面提示词', '负面提示词 V1'],
    ['第7轮', '校验优化 Agent', '去重、冲突检查、长度控制', '最终提示词'],
    ['第8轮', '结果输出 Agent', '输出完整优化结果包', '结果供用户确认'],
]
add_table(doc, flow_steps[0], flow_steps[1:], col_widths=[1.5, 3, 5, 3])

add_heading(doc, '3.2 用户体验设计', level=2)
add_bullet(doc, '加载动画：优化过程中显示"Hermes Agent 智能优化中..."，分阶段展示进度。')
add_bullet(doc, '结果对比：优化完成后以左右分栏展示原始描述与优化后提示词，高亮变化部分。')
add_bullet(doc, '一键采纳：用户可一键采纳优化结果，或手动编辑后生成。')
add_bullet(doc, '耗时控制：标准优化≤8秒，超时自动降级为快速优化模式。')

doc.add_page_break()

# ========== 四、主体创建与图片生成主流程（重点） ==========
add_heading(doc, '四、主体创建与图片生成主流程（重点）', level=1)

add_heading(doc, '4.1 主流程图', level=2)
add_para(doc, '图4-1  主体创建与图片生成主流程图', indent=False)
# 插入真正的主流程图
doc.add_picture(r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\main_flow_diagram.png', width=Inches(5.2))
last_paragraph = doc.paragraphs[-1]
last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

add_heading(doc, '4.2 流程关键节点说明', level=2)
add_bullet(doc, '自动识别风格与镜头：根据用户描述文本自动判断风格（写实/国风/动漫）和镜头类型，降低操作门槛。')
add_bullet(doc, 'Hermes Agent 8轮优化：核心环节，通过底层隐式对话完成提示词优化，注入剧情上下文，匹配Skills模板。')
add_bullet(doc, '用户确认：用户可查看优化结果，一键采纳或手动编辑，确保生成前用户明确认可。')
add_bullet(doc, '质量检测闭环：生成后自动检测四类问题，不达标时触发最多3轮迭代优化，形成自优化闭环。')

doc.add_page_break()

# ========== 五、Skills 模板库 ==========
add_heading(doc, '五、Skills 模板库概览', level=1)

skills_stats = [
    ['分类', '模板数量', '适用主体', '优化轮次'],
    ['角色基础设定', '12套', '人物角色', '6轮'],
    ['角色三视图设定', '3套', '人物角色', '5轮'],
    ['角色一致性锁定', '4套', '人物角色', '4轮'],
    ['室内场景', '22套', '场景', '5轮'],
    ['室外场景', '18套', '场景', '5轮'],
    ['道具设计', '15套', '道具', '4轮'],
    ['通用质量增强', '8套', '全部', '3轮'],
    ['风格迁移', '10套', '全部', '4轮'],
    ['合计', '92套', '-', '-'],
]
add_table(doc, skills_stats[0], skills_stats[1:], col_widths=[4, 2.5, 3, 2.5])

add_para(doc, '每个 Skills 模板包含：系统角色定义、多轮对话流程、输出格式规范、质量控制规则四个部分，参考短剧编剧三模板架构和视频生成分镜模板设计。')

doc.add_page_break()

# ========== 六、跨角色协作泳道设计（重点） ==========
add_heading(doc, '六、跨角色协作泳道设计（重点）', level=1)

add_heading(doc, '6.1 泳道图', level=2)
add_para(doc, '图6-1  提示词优化跨角色协作泳道图', indent=False)
# 插入真正的泳道图
doc.add_picture(r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\swimlane_diagram.png', width=Inches(6.2))
last_paragraph = doc.paragraphs[-1]
last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

add_heading(doc, '6.2 关键协作节点说明', level=2)
add_bullet(doc, '节点1（用户→前端）：用户输入主体描述并点击优化，前端立即显示加载动画，避免用户等待焦虑。')
add_bullet(doc, '节点2（前端→Hermes）：前端调用 Hermes Agent 接口，传递主体信息，Hermes 开始多轮隐式对话。')
add_bullet(doc, '节点3（Hermes→前端）：Hermes 完成8轮优化后返回结果，前端以对比卡片形式展示。')
add_bullet(doc, '节点4（用户确认）：用户确认优化结果，前端提交图片生成请求。')
add_bullet(doc, '节点5（质量检测→前端）：质量检测服务返回检测报告，达标则保存，不达标则自动触发 Hermes 迭代优化。')

doc.add_page_break()

# ========== 七、实施计划与风险 ==========
add_heading(doc, '七、实施计划与风险', level=1)

add_heading(doc, '7.1 项目排期', level=2)
schedule = [
    ['阶段', '任务', '工期'],
    ['第一阶段', '需求评审与技术方案确认', '1周'],
    ['第二阶段', 'Hermes Agent Skills 模板开发', '2周'],
    ['第三阶段', '前端主体设置页面升级', '2周'],
    ['第四阶段', '质量检测服务开发', '1.5周'],
    ['第五阶段', '集成测试与联调', '1周'],
    ['第六阶段', '灰度发布与反馈收集', '1周'],
    ['合计', '', '8.5周'],
]
add_table(doc, schedule[0], schedule[1:], col_widths=[3, 7, 2])

add_heading(doc, '7.2 核心风险与应对', level=2)
risks = [
    ['风险', '等级', '应对措施'],
    ['Hermes 响应时间过长', '高', '设置超时降级，优化 Skills 减少轮次'],
    ['优化效果不稳定', '高', 'A/B 测试机制，持续优化 Skills 模板'],
    ['用户不接受新交互', '中', '提供新旧模式切换开关，逐步过渡'],
    ['并发量过大', '中', '限流+排队机制，弹性扩容'],
]
add_table(doc, risks[0], risks[1:], col_widths=[4, 2, 6])

add_heading(doc, '7.3 预期收益', level=2)
add_bullet(doc, '用户体验：图片一次成功率从55%提升至85%，重生成次数减少57%。')
add_bullet(doc, '技术壁垒：建立 Hermes Agent 多轮对话提示词优化的核心技术能力。')
add_bullet(doc, '成本节约：减少用户重生成带来的算力消耗，预计降低30%图片生成成本。')
add_bullet(doc, '可扩展性：该架构可复用于视频生成提示词优化、分镜脚本优化等场景。')

# 保存文档
output_path = r'D:\phpstudy_pro\WWW\ai-novel-workstation1.2.0\人工智能短剧系统_图片生成提示词优化产品方案_HermesAgent领导汇报版V2_含专业图形.docx'
doc.save(output_path)
print(f'文档已生成：{output_path}')
