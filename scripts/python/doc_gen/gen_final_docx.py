import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn

def set_chinese_font(run):
    """设置中文字体为宋体"""
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def add_heading(doc, text, level):
    """添加标题"""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        set_chinese_font(run)
        if level == 1:
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif level == 2:
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
        else:
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)

def add_paragraph(doc, text, bold=False):
    """添加段落"""
    p = doc.add_paragraph()
    if text.startswith('- ') or text.startswith('* '):
        p.style = 'List Bullet'
        text = text[2:]
    elif any(text.startswith(f"{i}. ") for i in range(1, 100)):
        p.style = 'List Number'
        idx = text.find(". ") + 2
        text = text[idx:]
    
    run = p.add_run(text)
    set_chinese_font(run)
    run.font.size = Pt(12)
    if bold:
        run.font.bold = True
    return p

def parse_markdown(md_file, docx_file):
    """解析Markdown并生成DOCX"""
    doc = Document()
    
    # 全局设置中文字体
    doc.styles['Normal'].font.name = '宋体'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    doc.styles['Normal'].font.size = Pt(12)
    
    if not os.path.exists(md_file):
        print(f"Error: {md_file} not found")
        return

    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    for line in lines:
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue
        
        if line.startswith('# '):
            add_heading(doc, line[2:], 1)
        elif line.startswith('## '):
            add_heading(doc, line[3:], 2)
        elif line.startswith('### '):
            add_heading(doc, line[4:], 3)
        elif line.startswith('#### '):
            add_heading(doc, line[5:], 4)
        elif line.startswith('---'):
            doc.add_page_break()
        elif line.startswith('[图'):
            p = doc.add_paragraph()
            p.alignment = 1 # Center
            run = p.add_run(line)
            set_chinese_font(run)
            run.font.italic = True
            run.font.size = Pt(10)
        else:
            add_paragraph(doc, line)
            
    doc.save(docx_file)
    print(f"Successfully generated: {docx_file}")

if __name__ == "__main__":
    md_path = 'AI短剧平台用户操作使用说明书_Final.md'
    docx_path = 'AI短剧平台用户操作使用说明书_正式版.docx'
    parse_markdown(md_path, docx_path)
