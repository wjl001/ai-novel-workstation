
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_font(run, font_name='微软雅黑', size=None, bold=False, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    if size:
        run.font.size = Pt(size)
    if bold:
        run.bold = True
    if color:
        run.font.color.rgb = color

def add_heading_with_font(doc, text, level, font_name='微软雅黑'):
    heading = doc.add_heading('', level=level)
    run = heading.add_run(text)
    size = 18 if level == 1 else (16 if level == 2 else 14)
    set_font(run, font_name, size=size, bold=True)
    return heading

def create_report():
    doc = Document()
    
    # --- Title ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run('AI短剧PC端多模型切换功能产品汇报方案')
    set_font(title_run, size=24, bold=True)
    
    doc.add_paragraph() # Spacer
    
    # --- Introduction ---
    p = doc.add_paragraph()
    run = p.add_run('一、 文档制作规范与参考依据')
    set_font(run, size=16, bold=True)
    
    p = doc.add_paragraph()
    run = p.add_run('1.1 制作标准与规范')
    set_font(run, size=14, bold=True)
    
    standards = [
        '排版结构：严格参照服务器「软著参考资料目录」中《变量平台用户操作手册》《03-软件操作说明书》的整体框架。',
        '字体要求：全文统一采用微软雅黑字体，确保视觉呈现专业、严谨。',
        '目录结构：内置完整分级目录，便于管理层快速索引与评审。',
        '行文风格：采用企业正式书面语，逻辑清晰，通俗易懂，规避口语化表述。'
    ]
    for std in standards:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(std)
        set_font(run, size=11)
        
    p = doc.add_paragraph()
    run = p.add_run('1.2 参考资源路径')
    set_font(run, size=14, bold=True)
    
    refs = [
        '线上原型站点：http://172.25.24.86:5173/ai-short-drama-creator/new',
        '前端源码路径：D:\\phpstudy_pro\\WWW\\ai-novel-workstation2.7/src/views',
        '界面素材来源：基于桌面路径「AI短剧Pc端截图」文件夹内17张界面截图进行拆解与标注。'
    ]
    for ref in refs:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(ref)
        set_font(run, size=11)

    # --- Section 2 ---
    p = doc.add_paragraph()
    run = p.add_run('二、 核心设计基准（对标分镜导入闭环逻辑）')
    set_font(run, size=16, bold=True)
    
    p = doc.add_paragraph()
    run = p.add_run('本次多模型切换功能的逻辑设计，深度对标系统现有的「分镜导入」成熟模块。通过引入严谨的数据冲突校验与异常拦截机制，确保模型切换过程中的数据一致性与系统稳定性。')
    set_font(run, size=11)
    
    p = doc.add_paragraph()
    run = p.add_run('2.1 数据冲突处置机制')
    set_font(run, size=14, bold=True)
    
    p = doc.add_paragraph()
    run = p.add_run('当用户尝试切换已存在生成结果（图片/视频）的分镜模型时，系统将触发冲突处理逻辑：')
    set_font(run, size=11)
    
    conflicts = [
        '覆盖模式：系统将清空该分镜下的历史生成数据，使用新选定模型重新进行生成任务。此模式适用于用户对当前生成效果不满意，需全面替换风格的场景。',
        '跳过模式：系统保留该分镜已有的生成数据，模型切换仅对该分镜后续可能的重试生成生效。此模式适用于保护已有优质成果，仅针对未来任务进行调整的场景。'
    ]
    for conf in conflicts:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(conf)
        set_font(run, size=11)

    p = doc.add_paragraph()
    run = p.add_run('2.2 异常拦截与原子化操作')
    set_font(run, size=14, bold=True)
    
    p = doc.add_paragraph()
    run = p.add_run('参考分镜导入的格式校验逻辑，模型切换过程遵循「全量成功或全量失败」的原则：')
    set_font(run, size=11)
    
    exceptions = [
        '参数校验：强制校验所选模型是否支持当前剧集的画幅比例（如9:16）、分辨率等硬性参数。',
        '任务拦截：若批量切换过程中任一分镜校验未通过，系统将立即终止本次全量切换任务，不执行任何数据库写入操作。',
        '同步反馈：系统通过全局提示弹窗告知用户具体的拦截原因（如“模型A不支持16:9比例，任务已终止”），引导用户进行正确配置。'
    ]
    for exc in exceptions:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(exc)
        set_font(run, size=11)

    # --- Section 3 ---
    p = doc.add_paragraph()
    run = p.add_run('三、 功能设计细则')
    set_font(run, size=16, bold=True)
    
    p = doc.add_paragraph()
    run = p.add_run('3.1 页面布局与入口设计')
    set_font(run, size=14, bold=True)
    
    layouts = [
        '配置入口：在分镜工作台（StoryboardView）顶部工具栏新增「模型配置」入口，采用醒目的图标标识。',
        '模型画廊：点击进入后，以网格卡片形式展示可用模型，包含模型名称、核心优势（如“写实增强”、“光影细腻”）及代表性示例图。',
        '当前生效标识：在模型卡片右上角通过徽章（Badge）清晰标注当前剧集正在使用的模型。'
    ]
    for lay in layouts:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(lay)
        set_font(run, size=11)

    p = doc.add_paragraph()
    run = p.add_run('3.2 核心交互按钮')
    set_font(run, size=14, bold=True)
    
    buttons = [
        '「确认切换」：触发选中模型的应用逻辑，并根据冲突检测结果弹出二次确认框。',
        '「一键应用至全剧集」：快速将所选模型同步至当前作品的所有分集中，大幅提升批量处理效率。',
        '「效果对比预览」：支持用户在不真正切换模型的情况下，针对单个分镜生成低分预览图，辅助决策。'
    ]
    for btn in buttons:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(btn)
        set_font(run, size=11)

    p = doc.add_paragraph()
    run = p.add_run('3.3 弹窗交互规范')
    set_font(run, size=14, bold=True)
    
    popups = [
        '冲突确认弹窗：文案清晰告知风险（如“切换模型将导致已生成的X条分镜视频被覆盖”），提供明确的选项按钮。',
        '进度反馈弹窗：在大规模批量切换时，展示线性进度条，并实时显示已完成比例，缓解用户焦虑。',
        '异常拦截弹窗：采用高对比度警示色，突出显示错误原因，并提供「前往调整」的快捷链接。'
    ]
    for pop in popups:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(pop)
        set_font(run, size=11)

    # --- Section 4 ---
    p = doc.add_paragraph()
    run = p.add_run('四、 全业务流程逻辑（全场景覆盖）')
    set_font(run, size=16, bold=True)
    
    p = doc.add_paragraph()
    run = p.add_run('4.1 正常切换流')
    set_font(run, size=14, bold=True)
    p = doc.add_paragraph('用户选择新模型 -> 系统校验兼容性 -> 检测无冲突分镜 -> 执行数据库更新 -> 界面实时刷新状态。')
    set_font(p.runs[0], size=11)

    p = doc.add_paragraph()
    run = p.add_run('4.2 冲突处理流')
    set_font(run, size=14, bold=True)
    p = doc.add_paragraph('检测到分镜已存在结果 -> 弹出冲突策略选择框 -> 用户选择「覆盖」 -> 系统清空旧资产并入库新配置。')
    set_font(p.runs[0], size=11)

    p = doc.add_paragraph()
    run = p.add_run('4.3 异常拦截流')
    set_font(run, size=14, bold=True)
    p = doc.add_paragraph('校验发现模型与画幅不匹配 -> 立即锁死「确认」按钮 -> 浮窗提示原因 -> 拦截任何数据变更请求。')
    set_font(p.runs[0], size=11)

    # --- Conclusion ---
    p = doc.add_paragraph()
    run = p.add_run('五、 方案总结与核心优势')
    set_font(run, size=16, bold=True)
    
    advantages = [
        '功能稳定性：对标分镜导入逻辑，通过严谨的异常拦截确保系统不产生脏数据。',
        '操作风险可控：多级确认机制与原子化更新，保障了用户资产的安全性。',
        '业务闭环：从模型选择、兼容校验到冲突处置，实现了全场景的闭环设计。'
    ]
    for adv in advantages:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(adv)
        set_font(run, size=11)

    # --- Footer ---
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run('北京君禾世纪科技有限公司 产品部\n2026年7月7日')
    set_font(run, size=11)

    # Save
    file_path = 'multi_model_report.docx'
    doc.save(file_path)
    print(f"File saved to {os.path.abspath(file_path)}")

if __name__ == '__main__':
    create_report()
