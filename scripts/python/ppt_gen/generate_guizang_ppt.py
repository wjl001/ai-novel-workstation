import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt, Cm
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

def create_guizang_ppt():
    prs = Presentation()
    # Set to 16:9 aspect ratio
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Theme colors
    bg_color = RGBColor(15, 15, 20) # Very dark blue/grey
    accent_color = RGBColor(0, 240, 255) # Cyberpunk Cyan
    accent2_color = RGBColor(255, 0, 60) # Neon Red
    text_color = RGBColor(240, 240, 240)
    sub_text_color = RGBColor(160, 160, 170)
    
    # Helper function to set slide background
    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = bg_color
        
        # Add some industrial decorative shapes
        decor1 = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1)
        )
        decor1.fill.solid()
        decor1.fill.fore_color.rgb = accent_color
        decor1.line.fill.background()
        
        decor2 = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(13.0), Inches(0.5), Inches(0.1), Inches(1.5)
        )
        decor2.fill.solid()
        decor2.fill.fore_color.rgb = accent2_color
        decor2.line.fill.background()

    # 1. 封面
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # blank layout
    set_slide_background(slide)
    
    # Title
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11.333), Inches(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "我们的突围战：做AI短剧界的“超级工厂”"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = text_color
    p.font.name = "Microsoft YaHei"
    
    # Subtitle
    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(3.5), Inches(11.333), Inches(1))
    tf2 = txBox2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "不跟大厂抢流量，做包工头最爱的生产力工具"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(28)
    p2.font.color.rgb = accent_color
    p2.font.name = "Microsoft YaHei"
    
    # Info
    txBox3 = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(11.333), Inches(1.5))
    tf3 = txBox3.text_frame
    info_text = "汇报主题：PC端AI短剧生产平台竞品调研与立项规划\n汇报人：产品部\n汇报日期：2026年5月"
    for line in info_text.split('\n'):
        p3 = tf3.add_paragraph() if tf3.paragraphs[0].text else tf3.paragraphs[0]
        p3.text = line
        p3.alignment = PP_ALIGN.CENTER
        p3.font.size = Pt(18)
        p3.font.color.rgb = sub_text_color
        p3.font.name = "Microsoft YaHei"

    def create_content_slide(title, subtitle, bullets, highlight_index=-1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        set_slide_background(slide)
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(1))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(40)
        p.font.bold = True
        p.font.color.rgb = text_color
        p.font.name = "Microsoft YaHei"
        
        # Subtitle / Core point
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11), Inches(0.8))
        tf_sub = sub_box.text_frame
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = "核心观点 | " + subtitle
        p_sub.font.size = Pt(22)
        p_sub.font.color.rgb = accent_color
        p_sub.font.name = "Microsoft YaHei"
        
        # Content
        content_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.5), Inches(4))
        tf_c = content_box.text_frame
        tf_c.word_wrap = True
        for i, bullet in enumerate(bullets):
            p_c = tf_c.add_paragraph() if i > 0 else tf_c.paragraphs[0]
            p_c.text = "• " + bullet
            p_c.font.size = Pt(24)
            p_c.font.name = "Microsoft YaHei"
            if i == highlight_index:
                p_c.font.color.rgb = accent2_color
                p_c.font.bold = True
            else:
                p_c.font.color.rgb = text_color
            p_c.space_after = Pt(20)

    # 2. 市场现状
    create_content_slide(
        "[市场现状] 市场变天：横店没戏拍了，AI掀了桌子",
        "真人短剧正被淘汰，AI短剧已经占据绝对统治地位。",
        [
            "大盘数据：2026年Q1微短剧上线12.8万部，AI制作高达12.2万部。",
            "市场占比：AI短剧渗透率已超 95%！",
            "行业现状：传统剧组无戏可拍，算力（显卡）正在取代场地和群演。"
        ],
        1 # Highlight 95%
    )

    # 3. 核心驱动力
    create_content_slide(
        "[核心驱动力] 为什么大家都疯了？账算得太明白了",
        "成本砍掉90%，资本和制作团队必然全面倒向AI。",
        [
            "传统模式：场地+群演+灯光 = 50-100万成本，周期半个月。",
            "AI模式：3-5人+电脑+键盘 = 3-5万成本，周期1个星期。",
            "结论：资本是逐利的，十倍的利润差是市场剧变的唯一动力。"
        ],
        2
    )

    # 4. 用户痛点
    create_content_slide(
        "[用户痛点] 工作室每天都在骂什么？",
        "虽然用了AI，但现在的工具根本没法干“重体力活”。",
        [
            "痛点一：“换头怪”灾难。角色脸部无法固定，观众极易出戏。",
            "痛点二：“缝衣式”工作流。写剧本(GPT) -> 跑图(MJ) -> 生成视频(可灵) -> 剪辑(剪映)。",
            "现状总结：几千个镜头，切几十个文件夹，大半夜能把人逼疯。"
        ],
        2
    )

    # 5. 竞品分析
    create_content_slide(
        "[竞品分析] 扒一扒对手：大厂们在干嘛？",
        "大厂都在拼底层大模型（造发动机），没空做体验（造车壳）。",
        [
            "字节（小云雀/即梦）：Seedance 2.0 模型，懂导演思维，镜头连贯。",
            "快手（可灵）：Kling 3.0，物理引擎牛，质感像电影大片。",
            "阿里（快乐小马）：便宜又快，自带翻译，主打出海管饱。"
        ]
    )

    # 6. 战略洞察
    create_content_slide(
        "[战略洞察] 为什么大厂不做“手机端短剧APP”？",
        "大厂绝不内部打架，工具端不能和流量端抢饭碗。",
        [
            "定位不同：AI工具是“炒菜的锅”，抖音快手是“大饭店”。",
            "流量闭环：大厂巴不得大家用他们的工具生成视频，再发回他们主站变现。",
            "物理限制：手机屏幕太小，搓几十集、几百个转场的长剧，完全不现实。"
        ]
    )

    # 7. 我们的定位
    create_content_slide(
        "[我们的定位] 避开大厂锋芒，做“包工头”最爱的软件",
        "我们不造发动机，我们造最好的全自动流水线机床！",
        [
            "绝对不碰：烧钱的底层大模型、C端短剧流量分发。",
            "死死咬住：B端工作室的“生产效率”和“工程化管理”。",
            "目标用户：急着赚钱、手里攥着剧本、被零碎软件折磨的短剧包工头。"
        ],
        1
    )

    # 8. 杀手锏 1
    create_content_slide(
        "[杀手锏 1] 剧本一键拆解，把“作坊”变成“正规军”",
        "消灭乱七八糟的文件夹，建立正规的工程化目录树。",
        [
            "一键导入：十几万字剧本拖入，后台大模型瞬间解析。",
            "结构化拆解：自动生成 [集数 -> 场次 -> 镜头分镜 -> 角色台词]。",
            "左树右图：左侧点选镜头，右侧直接出画面生成面板。"
        ]
    )

    # 9. 杀手锏 2
    create_content_slide(
        "[杀手锏 2] “锁死”那张脸，打造短剧界“横店资产库”",
        "彻底解决“换头怪”痛点，让数字演员随叫随到。",
        [
            "开拍前捏脸：建立专属“数字演员休息室”。",
            "资产固化：死死锁住人脸特征、服装风格、AI配音音色。",
            "跨集调用：第一集到第三十集，一键“调用男主A”，角度随便换，脸绝对不崩。"
        ]
    )

    # 10. 杀手锏 3
    create_content_slide(
        "[杀手锏 3] “海纳百川”的智能调度器",
        "聚合各家API，用户只需提需求，我们负责跑腿。",
        [
            "痛点：买三家会员，开三个网页。",
            "解法：全面接入大厂开放API。要网感调字节，要特效调可灵，要便宜调阿里。",
            "本质：做一个“AI接口二道贩子”，用极致体验赚工作流的钱。"
        ],
        2
    )

    # 11. 商业模式
    create_content_slide(
        "[商业模式] 怎么赚钱？（搞钱才是硬道理）",
        "细水长流收月租，赚差价收云租金，出海风口卖增值。",
        [
            "SaaS高级会员：卖剧本拆解、资产锁、无限轨剪辑等核心功能。",
            "算力差价与云存储：拿大厂批发价赚差价；收工程文件的云端保存费。",
            "“一键出海”增值包：按次收费，提供“一键翻译+口型同步”。"
        ]
    )

    # 12. 下一步行动
    create_content_slide(
        "[下一步行动] 我们的 Action Plan",
        "小步快跑，快速验证，死磕前端体验。",
        [
            "启动MVP开发：先做“剧本拆解成目录树” + “简单接入一个视频接口”。",
            "引入种子用户：找2-3家痛苦的短剧小团队入伙免费测，照着痛点改。",
            "前端资源倾斜：HomeView.vue等交互必须像专业软件一样丝滑，拒绝网页卡顿感。"
        ],
        2
    )

    # 13. 总结愿景
    create_content_slide(
        "[总结愿景] 卖铲子的人，永远最赚钱",
        "在AI短剧的淘金热中，我们就是那把最好用的冲锋枪。",
        [
            "旧规矩全被打破，影视工业面临重构。",
            "降维打击市面上的“网页玩具”。",
            "只要咱们这把“冲锋枪”造得好，必定能在这个风口赚得盆满钵满！"
        ]
    )

    # 14. 结尾 Q & A
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    
    txBox = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(11.333), Inches(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Q & A"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(72)
    p.font.bold = True
    p.font.color.rgb = accent_color
    p.font.name = "Microsoft YaHei"
    
    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(4.5), Inches(11.333), Inches(1))
    tf2 = txBox2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "感谢聆听，请领导批评指正"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(28)
    p2.font.color.rgb = text_color
    p2.font.name = "Microsoft YaHei"

    prs.save('AI短剧平台竞品调研与立项规划.pptx')
    print("PPT generated successfully!")

if __name__ == '__main__':
    create_guizang_ppt()
