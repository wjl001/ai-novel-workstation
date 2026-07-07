#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将Markdown格式的操作说明书转换成DOCX文档
"""

import sys
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def set_chinese_font(run):
    """设置中文字体"""
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def add_heading(doc, text, level):
    """添加标题"""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        set_chinese_font(run)
        if level <= 2:
            run.font.bold = True
            run.font.size = Pt(16 if level == 1 else 14)
        else:
            run.font.size = Pt(12)


def add_paragraph(doc, text, bold=False):
    """添加段落"""
    p = doc.add_paragraph()
    if text.startswith('- ') or text.startswith('* '):
        # 列表项
        p.style = 'List Bullet'
        text = text[2:]
    elif text.startswith('1. ') or text.startswith('2. ') or text.startswith('3. ') or text.startswith('4. ') or text.startswith('5. ') or text.startswith('6. ') or text.startswith('7. ') or text.startswith('8. ') or text.startswith('9. '):
        p.style = 'List Number'
        text = text[3:]
    
    run = p.add_run(text)
    set_chinese_font(run)
    run.font.size = Pt(12)
    if bold:
        run.font.bold = True
    return p


def add_code_block(doc, text):
    """添加代码块"""
    p = doc.add_paragraph()
    p.style = 'Code'
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(10)


def parse_markdown(md_file, docx_file):
    """解析Markdown并生成DOCX"""
    doc = Document()
    
    # 设置默认字体
    doc.styles['Normal'].font.name = '宋体'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    doc.styles['Normal'].font.size = Pt(12)
    
    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    i = 0
    in_code_block = False
    code_content = []
    
    while i < len(lines):
        line = lines[i].rstrip('\n')
        i += 1
        
        # 代码块
        if line.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_content = []
            else:
                add_code_block(doc, '\n'.join(code_content))
                in_code_block = False
            continue
        
        if in_code_block:
            code_content.append(line)
            continue
        
        # 处理空行
        if not line.strip():
            doc.add_paragraph()
            continue
        
        # 处理标题
        if line.startswith('# '):
            add_heading(doc, line[2:], 1)
        elif line.startswith('## '):
            add_heading(doc, line[3:], 2)
        elif line.startswith('### '):
            add_heading(doc, line[4:], 3)
        elif line.startswith('#### '):
            add_heading(doc, line[5:], 4)
        elif line.startswith('##### '):
            add_heading(doc, line[6:], 5)
        elif line.startswith('---'):
            # 分隔线
            p = doc.add_paragraph()
            p.add_run('_' * 50)
        elif line.startswith('**') and line.endswith('**'):
            # 加粗文本
            add_paragraph(doc, line[2:-2], bold=True)
        elif line.startswith('- ') or line.startswith('* '):
            # 列表项
            add_paragraph(doc, line)
        elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. ') or line.startswith('4. ') or line.startswith('5. ') or line.startswith('6. ') or line.startswith('7. ') or line.startswith('8. ') or line.startswith('9. '):
            # 编号列表
            add_paragraph(doc, line)
        elif line.startswith('> '):
            # 引用
            p = doc.add_paragraph()
            p.style = 'Intense Quote'
            run = p.add_run(line[2:])
            set_chinese_font(run)
        else:
            # 普通段落
            add_paragraph(doc, line)
    
    doc.save(docx_file)
    print(f"已生成: {docx_file}")


def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(script_dir)
    
    md_file = os.path.join(project_root, 'docs', 'AI短剧创作平台操作说明书.md')
    docx_file = os.path.join(project_root, 'docs', 'AI短剧创作平台操作说明书.docx')
    
    if not os.path.exists(md_file):
        print(f"错误: 找不到Markdown文件 {md_file}")
        sys.exit(1)
    
    # 确保输出目录存在
    os.makedirs(os.path.dirname(docx_file), exist_ok=True)
    
    parse_markdown(md_file, docx_file)


if __name__ == '__main__':
    main()
